"""Read-only profiling for tabular files before Origin is activated."""

from __future__ import annotations

import csv
import math
import re
from pathlib import Path
from typing import Any, Iterable

from .utils.hashing import sha256_file


class DataInspectionError(ValueError):
    """The source file cannot be inspected under the requested contract."""


_INTEGER = re.compile(r"^[+-]?\d+$")
_FLOAT = re.compile(
    r"^[+-]?(?:\d+\.\d*|\d*\.\d+|\d+)(?:[eE][+-]?\d+)?$"
)


_sha256 = sha256_file


def _coerce_text(value: str) -> Any:
    stripped = value.strip()
    if stripped == "":
        return None
    if _INTEGER.fullmatch(stripped):
        try:
            return int(stripped)
        except ValueError:
            return value
    if _FLOAT.fullmatch(stripped):
        try:
            number = float(stripped)
            return number if math.isfinite(number) else value
        except ValueError:
            return value
    return value


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        return _coerce_text(value)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _read_delimited(path: Path) -> tuple[list[str], list[list[Any]]]:
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = [
            [_normalize_cell(value) for value in row]
            for row in csv.reader(handle, delimiter=delimiter)
        ]
    return ["Data"], rows


def _read_xlsx(path: Path, sheet_name: str | None) -> tuple[list[str], str, list[list[Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = list(workbook.sheetnames)
        selected = sheet_name or sheets[0]
        if selected not in sheets:
            raise DataInspectionError(f"Excel worksheet does not exist: {selected}")
        worksheet = workbook[selected]
        rows = [
            [_normalize_cell(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        return sheets, selected, rows
    finally:
        workbook.close()


def _read_xls(path: Path, sheet_name: str | None) -> tuple[list[str], str, list[list[Any]]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheets = workbook.sheet_names()
        selected = sheet_name or sheets[0]
        if selected not in sheets:
            raise DataInspectionError(f"Excel worksheet does not exist: {selected}")
        sheet = workbook.sheet_by_name(selected)
        rows = [
            [_normalize_cell(sheet.cell_value(row, column)) for column in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
        return sheets, selected, rows
    finally:
        workbook.release_resources()


def _has_header(rows: list[list[Any]]) -> bool:
    if len(rows) < 2:
        return False
    first = rows[0]
    second = rows[1]
    return any(isinstance(value, str) and value.strip() for value in first) and any(
        isinstance(value, (int, float)) and not isinstance(value, bool) for value in second
    )


def _rectangular(rows: Iterable[list[Any]]) -> list[list[Any]]:
    materialized = [list(row) for row in rows]
    width = max((len(row) for row in materialized), default=0)
    return [row + [None] * (width - len(row)) for row in materialized]


def _profile(label: str, values: list[Any], index: int) -> dict[str, Any]:
    non_empty = [value for value in values if value is not None]
    numeric = [
        value
        for value in non_empty
        if isinstance(value, (int, float)) and not isinstance(value, bool)
    ]
    text = [value for value in non_empty if isinstance(value, str)]
    if numeric and text:
        inferred = "mixed"
    elif numeric:
        inferred = "numeric"
    elif text:
        inferred = "text"
    else:
        inferred = "empty"
    return {
        "index": index,
        "label": label,
        "inferred_type": inferred,
        "non_empty_count": len(non_empty),
        "numeric_count": len(numeric),
        "text_count": len(text),
        "missing_count": len(values) - len(non_empty),
        "first_non_empty": non_empty[0] if non_empty else None,
        "last_non_empty": non_empty[-1] if non_empty else None,
        "unique_count": len({str(value) for value in non_empty}),
    }


def inspect_data_source(
    source_path: str | Path,
    *,
    sheet_name: str | None = None,
    has_header: bool | None = None,
) -> dict[str, Any]:
    """Inspect one local tabular source without opening or activating Origin."""

    path = Path(source_path).expanduser().resolve()
    if not path.is_file():
        raise DataInspectionError(f"Data source does not exist: {path}")
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        sheets, rows = _read_delimited(path)
        selected = "Data"
        if sheet_name not in {None, selected}:
            raise DataInspectionError(f"Delimited data has no worksheet named: {sheet_name}")
    elif suffix in {".xlsx", ".xlsm"}:
        sheets, selected, rows = _read_xlsx(path, sheet_name)
    elif suffix == ".xls":
        sheets, selected, rows = _read_xls(path, sheet_name)
    else:
        raise DataInspectionError(f"Unsupported data source extension: {suffix or '<none>'}")

    rows = _rectangular(rows)
    header = _has_header(rows) if has_header is None else bool(has_header)
    if header and rows:
        labels = [
            str(value).strip() if value not in {None, ""} else f"Column {index + 1}"
            for index, value in enumerate(rows[0])
        ]
        data = rows[1:]
    else:
        width = len(rows[0]) if rows else 0
        labels = [f"Column {index + 1}" for index in range(width)]
        data = rows
    data = _rectangular(data)
    profiles = [
        _profile(label, [row[index] for row in data], index)
        for index, label in enumerate(labels)
    ]
    duplicates = sorted({label for label in labels if labels.count(label) > 1})
    return {
        "source_path": str(path),
        "source_sha256": _sha256(path),
        "file_type": suffix.lstrip("."),
        "sheets": sheets,
        "selected_sheet": selected,
        "has_header": header,
        "rows": len(data),
        "columns": len(labels),
        "labels": labels,
        "duplicate_labels": duplicates,
        "column_profiles": profiles,
    }

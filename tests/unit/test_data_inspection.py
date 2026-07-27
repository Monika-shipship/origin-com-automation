import csv
from pathlib import Path

import openpyxl
import pytest

from origin_com_automation.data_inspection import DataInspectionError, inspect_data_source


def test_csv_inspection_profiles_numeric_text_mixed_and_missing_columns(tmp_path: Path):
    source = tmp_path / "mixed.csv"
    with source.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Lch", "Ion", "Rc", "Unit"])
        writer.writerow([0.041, 2260, "NA", "uA/um"])
        writer.writerow([0.080, 1800, "~1", "uA/um"])
        writer.writerow([0.120, 950, 0.42, "uA/um"])

    result = inspect_data_source(source, has_header=True)
    profiles = {item["label"]: item for item in result["column_profiles"]}

    assert result["file_type"] == "csv"
    assert result["rows"] == 3
    assert result["columns"] == 4
    assert len(result["source_sha256"]) == 64
    assert profiles["Ion"]["numeric_count"] == 3
    assert profiles["Ion"]["first_non_empty"] == 2260
    assert profiles["Ion"]["last_non_empty"] == 950
    assert profiles["Rc"]["numeric_count"] == 1
    assert profiles["Rc"]["text_count"] == 2
    assert profiles["Rc"]["inferred_type"] == "mixed"


def test_xlsx_inspection_selects_named_sheet_and_preserves_unicode_labels(tmp_path: Path):
    source = tmp_path / "benchmark.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["长度 (µm)", "Ion (µA/µm)", "Class"])
    sheet.append([0.1, 2260, "this work"])
    sheet.append([0.2, 950, None])
    workbook.create_sheet("Notes").append(["not data"])
    workbook.save(source)

    result = inspect_data_source(source, sheet_name="Data", has_header=True)

    assert result["sheets"] == ["Data", "Notes"]
    assert result["selected_sheet"] == "Data"
    assert result["labels"] == ["长度 (µm)", "Ion (µA/µm)", "Class"]
    assert result["column_profiles"][2]["missing_count"] == 1


def test_inspection_rejects_missing_sheet_and_unsupported_extension(tmp_path: Path):
    source = tmp_path / "data.xlsx"
    openpyxl.Workbook().save(source)

    with pytest.raises(DataInspectionError, match="worksheet"):
        inspect_data_source(source, sheet_name="Missing")

    unsupported = tmp_path / "data.json"
    unsupported.write_text("{}", encoding="utf-8")
    with pytest.raises(DataInspectionError, match="Unsupported"):
        inspect_data_source(unsupported)


def test_irregular_csv_detects_metadata_header_decimal_and_effective_range(tmp_path: Path):
    source = tmp_path / "irregular data.csv"
    source.write_text(
        "Instrument export;;;\n"
        "Operator note;;;\n"
        "Lch;Ion;Ion;Class\n"
        "0,041;2.260E3;2260;this work\n"
        "0,080;1.800E3;1800;literature\n"
        ";;;\n",
        encoding="utf-8-sig",
    )

    result = inspect_data_source(source)

    assert result["encoding"] == "utf-8-sig"
    assert result["delimiter"] == ";"
    assert result["decimal_convention"] == "comma"
    assert result["selected_header_row"] == 2
    assert result["metadata_rows"] == [
        ["Instrument export", None, None, None],
        ["Operator note", None, None, None],
    ]
    assert result["labels"] == ["Lch", "Ion", "Ion", "Class"]
    assert result["duplicate_labels"] == ["Ion"]
    assert result["effective_row_range"] == [3, 4]
    assert result["rows"] == 2
    assert result["preview_head"][0][:3] == [0.041, 2260.0, 2260]
    assert result["preview_tail"][-1][0] == 0.08
    assert result["header_candidates"][0]["row"] == 2


def test_ambiguous_header_is_reported_instead_of_hidden(tmp_path: Path):
    source = tmp_path / "ambiguous.csv"
    source.write_text("Name,Value\nLabel,Reading\nA,1\nB,2\n", encoding="utf-8")

    result = inspect_data_source(source)

    assert result["has_header"] is True
    assert result["ambiguities"]
    assert result["ambiguities"][0]["code"] == "HEADER_AMBIGUOUS"


def test_explicit_header_row_overrides_detection(tmp_path: Path):
    source = tmp_path / "explicit.csv"
    source.write_text("note,\nx,y\n1,2\n", encoding="utf-8")

    result = inspect_data_source(source, has_header=True, header_row=1)

    assert result["selected_header_row"] == 1
    assert result["labels"] == ["x", "y"]
    assert result["rows"] == 1

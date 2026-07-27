"""Read-only profiling for tabular files before Origin is activated."""

from __future__ import annotations

import csv
import io
import math
import re
from pathlib import Path
from typing import Any, Iterable

from .utils.hashing import sha256_file


class DataInspectionError(ValueError):
    """The source file cannot be inspected under the requested contract."""


_INTEGER = re.compile(r"^[+-]?\d+$")
_FLOAT = re.compile(
    r"^[+-]?(?:\d+\.\d*|\d*\.\d+|\d+)(?:[eE][+-]?\d+)?$"
)
_COMMA_FLOAT = re.compile(
    r"^[+-]?(?:\d+,\d*|\d*,\d+)(?:[eE][+-]?\d+)?$"
)


_sha256 = sha256_file


def _coerce_text(value: str, *, decimal_convention: str | None = None) -> Any:
    stripped = value.strip()
    if stripped == "":
        return None
    if _INTEGER.fullmatch(stripped):
        try:
            return int(stripped)
        except ValueError:
            return value
    if _FLOAT.fullmatch(stripped):
        try:
            number = float(stripped)
            return number if math.isfinite(number) else value
        except ValueError:
            return value
    if decimal_convention == "comma" and _COMMA_FLOAT.fullmatch(stripped):
        try:
            number = float(stripped.replace(",", "."))
            return number if math.isfinite(number) else value
        except ValueError:
            return value
    return value


def _normalize_cell(value: Any, *, decimal_convention: str | None = None) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        return _coerce_text(value, decimal_convention=decimal_convention)
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _decode_delimited(path: Path, requested: str | None) -> tuple[str, str]:
    encodings = [requested] if requested else ["utf-8-sig", "utf-16", "gb18030", "cp1252"]
    failures: list[str] = []
    for encoding in encodings:
        if encoding is None:
            continue
        try:
            return path.read_text(encoding=encoding), encoding
        except (UnicodeDecodeError, UnicodeError) as exc:
            failures.append(f"{encoding}: {exc}")
    raise DataInspectionError(
        "Could not decode delimited source with the allowed encodings: "
        + "; ".join(failures)
    )


def _sniff_delimiter(text: str, path: Path, requested: str | None) -> str:
    if requested is not None:
        if len(requested) != 1:
            raise DataInspectionError("delimiter must be exactly one character")
        return requested
    if path.suffix.lower() == ".tsv":
        return "\t"
    sample = text[:65536]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def _decimal_convention(
    rows: list[list[Any]], delimiter: str | None, requested: str | None
) -> str | None:
    if requested is not None:
        if requested not in {"dot", "comma"}:
            raise DataInspectionError("decimal_convention must be dot or comma")
        return requested
    values = [str(value).strip() for row in rows for value in row if str(value).strip()]
    comma = sum(bool(_COMMA_FLOAT.fullmatch(value)) for value in values)
    dot = sum(bool(_FLOAT.fullmatch(value)) and "." in value for value in values)
    if comma and (comma > dot or delimiter != ","):
        return "comma"
    if dot:
        return "dot"
    return None


def _read_delimited(
    path: Path,
    *,
    encoding: str | None,
    delimiter: str | None,
    decimal_convention: str | None,
) -> tuple[list[str], list[list[Any]], str, str, str | None]:
    text, selected_encoding = _decode_delimited(path, encoding)
    selected_delimiter = _sniff_delimiter(text, path, delimiter)
    raw_rows = [list(row) for row in csv.reader(io.StringIO(text), delimiter=selected_delimiter)]
    selected_decimal = _decimal_convention(
        raw_rows, selected_delimiter, decimal_convention
    )
    rows = [
        [
            _normalize_cell(value, decimal_convention=selected_decimal)
            for value in row
        ]
        for row in raw_rows
    ]
    return ["Data"], rows, selected_encoding, selected_delimiter, selected_decimal


def _read_xlsx(path: Path, sheet_name: str | None) -> tuple[list[str], str, list[list[Any]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = list(workbook.sheetnames)
        selected = sheet_name or sheets[0]
        if selected not in sheets:
            raise DataInspectionError(f"Excel worksheet does not exist: {selected}")
        worksheet = workbook[selected]
        rows = [
            [_normalize_cell(value) for value in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
        return sheets, selected, rows
    finally:
        workbook.close()


def _read_xls(path: Path, sheet_name: str | None) -> tuple[list[str], str, list[list[Any]]]:
    import xlrd

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheets = workbook.sheet_names()
        selected = sheet_name or sheets[0]
        if selected not in sheets:
            raise DataInspectionError(f"Excel worksheet does not exist: {selected}")
        sheet = workbook.sheet_by_name(selected)
        rows = [
            [_normalize_cell(sheet.cell_value(row, column)) for column in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]
        return sheets, selected, rows
    finally:
        workbook.release_resources()


def _has_header(rows: list[list[Any]]) -> bool:
    if len(rows) < 2:
        return False
    first = rows[0]
    second = rows[1]
    return any(isinstance(value, str) and value.strip() for value in first) and any(
        isinstance(value, (int, float)) and not isinstance(value, bool) for value in second
    )


def _empty_row(row: list[Any]) -> bool:
    return not any(value is not None and str(value).strip() for value in row)


def _trim_trailing_empty(rows: list[list[Any]]) -> list[list[Any]]:
    trimmed = list(rows)
    while trimmed and _empty_row(trimmed[-1]):
        trimmed.pop()
    return trimmed


def _header_candidates(rows: list[list[Any]]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for index, row in enumerate(rows[:25]):
        non_empty = [value for value in row if value is not None and str(value).strip()]
        if not non_empty:
            continue
        text_count = sum(isinstance(value, str) for value in non_empty)
        unique_count = len({str(value).strip().casefold() for value in non_empty})
        following = rows[index + 1:index + 6]
        following_values = [
            value
            for next_row in following
            for value in next_row
            if value is not None and str(value).strip()
        ]
        following_numeric = sum(
            isinstance(value, (int, float)) and not isinstance(value, bool)
            for value in following_values
        )
        transition = following_numeric / len(following_values) if following_values else 0.0
        density = len(non_empty) / max(len(row), 1)
        score = (
            float(text_count)
            + float(unique_count) / max(len(non_empty), 1)
            + 3.0 * transition
            + density
            - (2.0 if len(non_empty) == 1 else 0.0)
        )
        candidates.append(
            {
                "row": index,
                "score": round(score, 6),
                "non_empty_count": len(non_empty),
                "text_count": text_count,
                "following_numeric_ratio": round(transition, 6),
            }
        )
    return sorted(candidates, key=lambda item: (-item["score"], item["row"]))


def _rectangular(rows: Iterable[list[Any]]) -> list[list[Any]]:
    materialized = [list(row) for row in rows]
    width = max((len(row) for row in materialized), default=0)
    return [row + [None] * (width - len(row)) for row in materialized]


def _profile(label: str, values: list[Any], index: int) -> dict[str, Any]:
    non_empty = [value for value in values if value is not None]
    numeric = [
        value
        for value in non_empty
        if isinstance(value, (int, float)) and not isinstance(value, bool)
    ]
    text = [value for value in non_empty if isinstance(value, str)]
    if numeric and text:
        inferred = "mixed"
    elif numeric:
        inferred = "numeric"
    elif text:
        inferred = "text"
    else:
        inferred = "empty"
    return {
        "index": index,
        "label": label,
        "inferred_type": inferred,
        "non_empty_count": len(non_empty),
        "numeric_count": len(numeric),
        "text_count": len(text),
        "missing_count": len(values) - len(non_empty),
        "first_non_empty": non_empty[0] if non_empty else None,
        "last_non_empty": non_empty[-1] if non_empty else None,
        "unique_count": len({str(value) for value in non_empty}),
    }


def inspect_data_source(
    source_path: str | Path,
    *,
    sheet_name: str | None = None,
    has_header: bool | None = None,
    header_row: int | None = None,
    encoding: str | None = None,
    delimiter: str | None = None,
    decimal_convention: str | None = None,
    preview_rows: int = 3,
) -> dict[str, Any]:
    """Inspect one local tabular source without opening or activating Origin."""

    path = Path(source_path).expanduser().resolve()
    if not path.is_file():
        raise DataInspectionError(f"Data source does not exist: {path}")
    suffix = path.suffix.lower()
    if preview_rows < 1 or preview_rows > 25:
        raise DataInspectionError("preview_rows must be between 1 and 25")
    selected_encoding: str | None = None
    selected_delimiter: str | None = None
    selected_decimal: str | None = None
    if suffix in {".csv", ".tsv"}:
        sheets, rows, selected_encoding, selected_delimiter, selected_decimal = (
            _read_delimited(
                path,
                encoding=encoding,
                delimiter=delimiter,
                decimal_convention=decimal_convention,
            )
        )
        selected = "Data"
        if sheet_name not in {None, selected}:
            raise DataInspectionError(f"Delimited data has no worksheet named: {sheet_name}")
    elif suffix in {".xlsx", ".xlsm"}:
        sheets, selected, rows = _read_xlsx(path, sheet_name)
    elif suffix == ".xls":
        sheets, selected, rows = _read_xls(path, sheet_name)
    else:
        raise DataInspectionError(f"Unsupported data source extension: {suffix or '<none>'}")

    rows = _rectangular(_trim_trailing_empty(rows))
    if header_row is not None and has_header is False:
        raise DataInspectionError("header_row cannot be used when has_header is false")
    if header_row is not None and (header_row < 0 or header_row >= len(rows)):
        raise DataInspectionError("header_row is outside the source row range")
    candidates = _header_candidates(rows)
    ambiguities: list[dict[str, Any]] = []
    if has_header is False:
        header = False
        selected_header_row = None
    elif header_row is not None:
        header = True
        selected_header_row = header_row
    elif has_header is True:
        header = True
        selected_header_row = 0
    else:
        best = candidates[0] if candidates else None
        header = bool(best and best["text_count"] > 0 and best["score"] >= 3.0)
        selected_header_row = int(best["row"]) if header and best else None
        if header and len(candidates) > 1:
            second = candidates[1]
            if best is not None and best["score"] - second["score"] <= 1.25:
                ambiguities.append(
                    {
                        "code": "HEADER_AMBIGUOUS",
                        "message": "Multiple rows are plausible headers; set header_row explicitly",
                        "candidates": [best, second],
                    }
                )
    metadata_rows = rows[:selected_header_row] if selected_header_row is not None else []
    if header and rows:
        assert selected_header_row is not None
        header_values = rows[selected_header_row]
        labels = [
            str(value).strip() if value not in {None, ""} else f"Column {index + 1}"
            for index, value in enumerate(header_values)
        ]
        data_start = selected_header_row + 1
        data = rows[data_start:]
    else:
        width = len(rows[0]) if rows else 0
        labels = [f"Column {index + 1}" for index in range(width)]
        data_start = 0
        data = rows
    data = _rectangular(data)
    profiles = [
        _profile(label, [row[index] for row in data], index)
        for index, label in enumerate(labels)
    ]
    duplicates = sorted({label for label in labels if labels.count(label) > 1})
    return {
        "source_path": str(path),
        "source_sha256": _sha256(path),
        "file_type": suffix.lstrip("."),
        "sheets": sheets,
        "selected_sheet": selected,
        "encoding": selected_encoding,
        "delimiter": selected_delimiter,
        "decimal_convention": selected_decimal,
        "has_header": header,
        "selected_header_row": selected_header_row,
        "header_candidates": candidates,
        "ambiguities": ambiguities,
        "metadata_rows": metadata_rows,
        "effective_row_range": (
            [data_start, data_start + len(data) - 1] if data else None
        ),
        "rows": len(data),
        "columns": len(labels),
        "labels": labels,
        "duplicate_labels": duplicates,
        "column_profiles": profiles,
        "preview_head": data[:preview_rows],
        "preview_tail": data[-preview_rows:] if data else [],
    }

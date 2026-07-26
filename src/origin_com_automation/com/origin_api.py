"""Origin 2024 COM adapter with ownership and thread-boundary safeguards."""

from __future__ import annotations

import csv
import inspect
import json
import logging
import math
import os
import re
import shutil
import subprocess
import tempfile
import time
from collections import Counter
from collections.abc import Callable, Iterable, Mapping
from dataclasses import replace
from functools import wraps
from pathlib import Path
from threading import RLock
from typing import Any
from uuid import uuid4

from ..contracts import Artifact, ResultEnvelope
from ..config import PluginConfig
from ..native.common import NativeValidationError, RangeRef
from ..native.operations import (
    AnalysisOperationRegistry,
    build_analysis_template_plan,
    execute_analysis_template_plan,
    execute_xfunction_plan,
    read_analysis_operation,
    recalculate_operation,
)
from ..native.xfunctions import build_xfunction_plan
from ..objects.connectors import build_connector_plan
from ..objects.images import build_image_plan
from ..objects.matrices import build_matrix_plan
from ..objects.validation import ObjectPlanError
from ..objects.worksheets import TransformValidationError, transform_table
from ..objects.project import ProjectObjectError, build_folder_plan, build_note_plan
from ..graphs.catalog import GraphCatalogError, validate_graph_request
from ..graphs.layout import GraphLayoutError, build_layout_plan
from ..graphs.preview import inspect_png
from ..graphs.templates import GraphTemplateError, apply_template_plan
from ..services.analysis import build_analysis_request, run_analysis_data
from ..services.exports import (
    export_suffixes,
    normalize_export_format,
    select_changed_artifact,
    snapshot_artifacts,
    validate_export_path,
)
from ..services.graphs import build_graph_spec
from ..services.projects import (
    SourceOverwriteError,
    copy_project,
    project_signature,
    validate_project_artifact,
)
from ..services.worksheets import table_from_com_value
from .discovery import discover_registrations, executable_version
from .errors import (
    AnalysisExecutionError,
    LabTalkExecutionError,
    NoActiveSessionError,
    NoExistingOriginError,
    OriginAutomationError,
    OriginTimeoutError,
    OwnershipUnverifiedError,
    ProjectCloseUnconfirmedError,
    ProjectSaveUnconfirmedError,
)
from .session import SessionManager
from .worker import SerialComWorker

ORIGIN_PAGE_WORKSHEET = 2
ORIGIN_PAGE_GRAPH = 3
ORIGIN_PLOT_LINE = 200
ORIGIN_PLOT_SCATTER = 201
ORIGIN_PLOT_LINESYMB = 202
ORIGIN_PLOT_COLUMN = 203
ORIGIN_CREATE_HIDDEN = 3
ORIGIN_ARRAY2D_VARIANT = 0
ORIGIN_ARRAY1D_STR = 8
ORIGIN_ARRAY2D_NUMERIC = 2
ORIGIN_ARRAY2D_TEXT_FULL_PRECISION = 4
ORIGIN_DATA_FORMAT_NUMERIC = 0
ORIGIN_DATA_FORMAT_TEXT = 1
ORIGIN_DATA_FORMAT_TEXT_NUMERIC = 9
ORIGIN_MISSING_VALUE = -1.23456789e-300

logger = logging.getLogger(__name__)

GRAPH_OPTION_KEYS = {
    "x_min",
    "x_max",
    "y_min",
    "y_max",
    "x_tick_step",
    "y_tick_step",
    "x_scale",
    "y_scale",
    "x_title",
    "y_title",
    "legend",
    "rescale",
    "plot_styles",
    "data_binding",
    "categorical_style",
}
SCALE_TYPES = {"linear": 0, "log10": 2, "ln": 8, "log2": 9}
WORKSHEET_DATA_FORMATS = {
    "auto": ORIGIN_ARRAY2D_VARIANT,
    "numeric": ORIGIN_ARRAY2D_NUMERIC,
    "string": ORIGIN_ARRAY2D_TEXT_FULL_PRECISION,
    "variant": ORIGIN_ARRAY2D_VARIANT,
}
SINGLE_INSTANCE_PROGIDS = {"Origin.ApplicationSI", "Origin.ApplicationCOMSI"}
OWNED_INSTANCE_PROGID = "Origin.Application"
ORIGIN_DISABLE_SAVE_PROMPT_SCRIPT = "doc -s;"
ORIGIN_DELAYED_EXIT_SCRIPT = "def timerproc { exit; } timer 1;"
LINE_CONNECTIONS = {
    "none": 0,
    "straight": 1,
    "two_point_segment": 2,
    "three_point_segment": 3,
    "b_spline": 8,
    "spline": 9,
    "step_horizontal": 11,
    "step_vertical": 12,
    "step_horizontal_center": 13,
    "step_vertical_center": 14,
    "bezier": 15,
}
DATA_BINDING_PLOT_TYPES = {
    "line": ORIGIN_PLOT_LINE,
    "scatter": ORIGIN_PLOT_SCATTER,
    "line_symbol": ORIGIN_PLOT_LINESYMB,
    "bar": ORIGIN_PLOT_COLUMN,
}
CATEGORY_SHAPES = {
    "square": 1,
    "circle": 2,
    "triangle_up": 3,
    "triangle_down": 4,
    "diamond": 5,
    "hexagon": 6,
    "star": 7,
    "cross": 8,
    "x": 9,
}
CATEGORY_FILLS = {"solid": 0, "open": 1, "hollow": 3}
ORIGIN_PLOT_DESIGNATION_Y = 0
ORIGIN_PLOT_DESIGNATION_X = 3
ORIGIN_PLOT_DESIGNATION_LABEL = 4
HRESULT_ERROR_CODES = {
    -2146959355: "CO_E_SERVER_EXEC_FAILURE",
    -2147418111: "RPC_CALL_REJECTED",
    -2147417848: "RPC_DISCONNECTED",
    -2147023174: "RPC_SERVER_UNAVAILABLE",
}
RETRYABLE_READ_ERRORS = {"RPC_CALL_REJECTED", "RPC_SERVER_UNAVAILABLE"}
FILTER_OPERATORS: dict[str, Callable[[float, float], bool]] = {
    "gt": lambda value, threshold: value > threshold,
    "ge": lambda value, threshold: value >= threshold,
    "lt": lambda value, threshold: value < threshold,
    "le": lambda value, threshold: value <= threshold,
    "eq": lambda value, threshold: value == threshold,
    "ne": lambda value, threshold: value != threshold,
}


def _default_dispatch(progid: str, attach: bool) -> Any:
    import win32com.client

    if attach:
        # SI/COMSI decide whether to reuse an existing instance. PID auditing
        # after activation determines whether that instance is plugin-owned.
        return win32com.client.Dispatch(progid)
    # Origin.Application is documented to always create a new instance.
    return win32com.client.DispatchEx(progid)


def is_origin_process_name(name: str) -> bool:
    return Path(name).stem.lower() == "origin64"


def _default_process_snapshot() -> set[int]:
    try:
        import psutil

        return {
            int(process.info["pid"])
            for process in psutil.process_iter(["pid", "name"])
            if is_origin_process_name(str(process.info.get("name", "")))
        }
    except Exception:
        try:
            completed = subprocess.run(
                ["tasklist", "/FI", "IMAGENAME eq Origin64.exe", "/FO", "CSV", "/NH"],
                capture_output=True,
                text=True,
                check=False,
                timeout=3,
            )
            result: set[int] = set()
            for row in csv.reader(completed.stdout.splitlines()):
                if len(row) >= 2 and row[0].lower() == "origin64.exe":
                    result.add(int(row[1]))
            return result
        except Exception:
            return set()


def _safe_attr(obj: Any, name: str, default: Any = None) -> Any:
    try:
        value = getattr(obj, name)
        return value() if callable(value) and name in {"Count"} else value
    except Exception:
        return default


def _safe_call(obj: Any, name: str, *args: Any, default: Any = None) -> Any:
    try:
        method = getattr(obj, name)
        if not callable(method):
            return default
        return method(*args)
    except Exception:
        return default


def _collection_items(collection: Any) -> list[Any]:
    if collection is None:
        return []
    count = _safe_attr(collection, "Count")
    if count is not None:
        items: list[Any] = []
        for index in range(int(count)):
            try:
                item = collection.Item(index)
            except Exception:
                try:
                    item = collection(index)
                except Exception:
                    continue
            if item is not None:
                items.append(item)
        return items
    try:
        return list(collection)
    except Exception:
        return []


def _find_collection_item(collection: Any, reference: str) -> Any | None:
    if collection is None:
        return None
    try:
        item = collection.Item(reference)
        if item is not None:
            return item
    except Exception:
        pass
    normalized = reference.strip().casefold()
    for item in _collection_items(collection):
        names = {
            str(_safe_attr(item, "Name", "")).strip().casefold(),
            str(_safe_attr(item, "LongName", "")).strip().casefold(),
        }
        if normalized in names:
            return item
    return None


def _page_numeric_property(page: Any, name: str) -> float | None:
    value = _safe_attr(page, name)
    if value is None:
        value = _safe_call(page, "GetNumProp", name, default=None)
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) and number > 0 else None


def _matrix_data_object(matrix_sheet: Any) -> Any:
    for collection_name in ("MatrixObjects", "DataObjectBases"):
        objects = _collection_items(_safe_attr(matrix_sheet, collection_name))
        if objects:
            return objects[0]
    return matrix_sheet


def _matrix_object_block(
    matrix_object: Any,
    row: int,
    column: int,
    rows: int,
    columns: int,
) -> Any:
    args = (row, column, row + rows - 1, column + columns - 1)
    try:
        return matrix_object.GetData(*args, ORIGIN_ARRAY2D_VARIANT)
    except TypeError:
        return matrix_object.GetData(*args)


def _normalize_matrix_object_table(
    value: Any, *, column_major_bridge: bool
) -> list[list[Any]]:
    table = table_from_com_value(value)
    if column_major_bridge and table:
        table = [list(row) for row in zip(*table, strict=True)]
    return [[_normalize_cell(cell) for cell in row] for row in table]


def _project_path_parts(path: str) -> list[str]:
    return [part for part in path.replace("\\", "/").split("/") if part]


def _resolve_root_folder(root: Any, path: str) -> Any | None:
    current = root
    for part in _project_path_parts(path):
        current = _find_collection_item(_safe_attr(current, "Folders"), part)
        if current is None:
            return None
    return current


def _ensure_root_folder(root: Any, path: str) -> Any | None:
    current = root
    for part in _project_path_parts(path):
        folders = _safe_attr(current, "Folders")
        child = _find_collection_item(folders, part)
        if child is None:
            child = _safe_call(folders, "Add", part, default=None)
        if child is None:
            return None
        current = child
    return current


def _same_origin_object(left: Any, right: Any) -> bool:
    if left is right:
        return True
    try:
        return bool(left == right)
    except Exception:
        return False


def _required_attr(obj: Any, name: str, default: Any = None) -> Any:
    try:
        value = getattr(obj, name)
    except AttributeError:
        return default
    return value() if callable(value) and name == "Count" else value


def _required_collection_items(collection: Any) -> list[Any]:
    if collection is None:
        return []
    count = _required_attr(collection, "Count")
    if count is None:
        return list(collection)
    try:
        item_method = getattr(collection, "Item")
    except AttributeError:
        item_method = None
    items: list[Any] = []
    for index in range(int(count)):
        if callable(item_method):
            item = item_method(index)
        elif callable(collection):
            item = collection(index)
        else:
            raise TypeError("Origin COM collection exposes Count but no Item accessor")
        if item is not None:
            items.append(item)
    return items


def _labtalk_quote(value: str | Path) -> str:
    text = str(value).replace("\\", "/").replace('"', '\\"')
    return f'"{text}"'


def _same_file_path(left: Path, right: Path) -> bool:
    try:
        return os.path.samefile(left, right)
    except OSError:
        return left.resolve() == right.resolve()


def _file_identity(path: Path) -> tuple[int, int] | None:
    try:
        stat = path.stat()
    except OSError:
        return None
    return int(stat.st_dev), int(stat.st_ino)


def _is_optional_com_argument_error(exc: Exception) -> bool:
    hresult = getattr(exc, "hresult", None)
    if hresult is None and exc.args and isinstance(exc.args[0], int):
        hresult = exc.args[0]
    if isinstance(hresult, int) and hresult > 0x7FFFFFFF:
        hresult -= 0x100000000
    return hresult in {-2147352572, -2147352571, -2147352562}


def _column_string_values(column: Any, r1: int, r2: int) -> list[Any]:
    try:
        value = column.GetData(ORIGIN_ARRAY1D_STR, r1, r2)
    except Exception as exc:
        if not _is_optional_com_argument_error(exc):
            raise
        value = column.GetData(ORIGIN_ARRAY1D_STR, r1, r2, 0)
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def _is_missing_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value == ""
    if isinstance(value, float):
        if math.isnan(value):
            return True
        return math.isclose(value, ORIGIN_MISSING_VALUE, rel_tol=0.0, abs_tol=1e-315)
    return False


def _normalize_cell(value: Any) -> Any:
    return None if _is_missing_cell(value) else value


def _column_variant_values(column: Any, r1: int, r2: int) -> list[Any]:
    try:
        raw = column.GetData(ORIGIN_ARRAY2D_VARIANT, r1, r2)
    except Exception as exc:
        if not _is_optional_com_argument_error(exc):
            raise
        raw = column.GetData(ORIGIN_ARRAY2D_VARIANT, r1, r2, 0)
    if raw is None:
        return []
    if isinstance(raw, (list, tuple)):
        if all(not isinstance(item, (list, tuple)) for item in raw):
            return [_normalize_cell(item) for item in raw]
        values: list[Any] = []
        for item in raw:
            if isinstance(item, (list, tuple)):
                values.append(_normalize_cell(item[0] if item else None))
            else:
                values.append(_normalize_cell(item))
        return values
    return [_normalize_cell(raw)]


def _column_profile(values: Iterable[Any]) -> dict[str, int]:
    normalized = [_normalize_cell(value) for value in values]
    numeric_count = sum(
        isinstance(value, (int, float)) and not isinstance(value, bool)
        for value in normalized
        if value is not None
    )
    text_count = sum(isinstance(value, str) for value in normalized if value is not None)
    non_empty_count = sum(value is not None for value in normalized)
    return {
        "non_empty_count": non_empty_count,
        "numeric_count": numeric_count,
        "text_count": text_count,
        "missing_count": len(normalized) - non_empty_count,
    }


def _column_data_format(values: Iterable[Any]) -> int:
    profile = _column_profile(values)
    if profile["text_count"] and profile["numeric_count"]:
        return ORIGIN_DATA_FORMAT_TEXT_NUMERIC
    if profile["text_count"]:
        return ORIGIN_DATA_FORMAT_TEXT
    return ORIGIN_DATA_FORMAT_NUMERIC


def _set_column_data_format(column: Any, values: list[Any], *, force: bool) -> int | None:
    requested = _column_data_format(values)
    current = _safe_attr(column, "DataFormat")
    # A partial write containing text must not leave the column numeric. For
    # numeric-only partial writes, preserve an existing mixed/text format.
    should_set = force or requested in {
        ORIGIN_DATA_FORMAT_TEXT,
        ORIGIN_DATA_FORMAT_TEXT_NUMERIC,
    }
    if should_set:
        column.DataFormat = requested
        if requested == ORIGIN_DATA_FORMAT_TEXT_NUMERIC:
            try:
                column.TextAndNumericSetAlwaysAsText = False
            except Exception:
                pass
        return requested
    return int(current) if current is not None else None


def _rectangular_values(values: list[list[Any]]) -> tuple[list[list[Any]], int]:
    width = max((len(row) for row in values), default=0)
    return [list(row) + [None] * (width - len(row)) for row in values], width


def _cells_equal(expected: Any, actual: Any) -> bool:
    expected = _normalize_cell(expected)
    actual = _normalize_cell(actual)
    if expected is None or actual is None:
        return expected is actual
    if (
        isinstance(expected, (int, float))
        and not isinstance(expected, bool)
        and isinstance(actual, (int, float))
        and not isinstance(actual, bool)
    ):
        return math.isclose(float(expected), float(actual), rel_tol=1e-12, abs_tol=1e-12)
    return str(expected) == str(actual)


def _readback_matrix(
    sheet: Any,
    *,
    row: int,
    column: int,
    rows: int,
    columns: int,
) -> list[list[Any]]:
    if rows == 0 or columns == 0:
        return []
    raw = sheet.GetData(
        row,
        column,
        row + rows - 1,
        column + columns - 1,
        ORIGIN_ARRAY2D_VARIANT,
    )
    table = table_from_com_value(raw)
    return [
        [
            _normalize_cell(table[r][c])
            if r < len(table) and c < len(table[r])
            else None
            for c in range(columns)
        ]
        for r in range(rows)
    ]


def _matrix_mismatches(
    expected: list[list[Any]], actual: list[list[Any]], *, limit: int = 10
) -> list[dict[str, Any]]:
    mismatches: list[dict[str, Any]] = []
    for row_index, row in enumerate(expected):
        for column_index, expected_value in enumerate(row):
            actual_value = (
                actual[row_index][column_index]
                if row_index < len(actual) and column_index < len(actual[row_index])
                else None
            )
            if not _cells_equal(expected_value, actual_value):
                mismatches.append(
                    {
                        "row_offset": row_index,
                        "column_offset": column_index,
                        "expected": expected_value,
                        "actual": actual_value,
                    }
                )
                if len(mismatches) >= limit:
                    return mismatches
    return mismatches


def _system_worksheet_template(app: Any) -> tuple[str, bool]:
    """Prefer Origin's installation template over a user-customized Origin.otwu."""

    try:
        program_path = Path(str(app.Path(4))).expanduser().resolve()
    except Exception:
        return "Origin", False
    localization = program_path / "Localization"
    candidates = [
        localization / "E" / "ORIGIN.otwu",
        localization / "C" / "ORIGIN.otwu",
        program_path / "Templates" / "ORIGIN.otwu",
    ]
    if localization.is_dir():
        candidates.extend(sorted(localization.glob("*/ORIGIN.otwu")))
    for candidate in candidates:
        if candidate.is_file():
            return str(candidate), True
    return "Origin", False


def _project_structure_fingerprint(app: Any) -> dict[str, list[dict[str, Any]]]:
    fingerprint: dict[str, list[dict[str, Any]]] = {}
    worksheet_catalog = _worksheet_column_catalog(app)
    for collection_name in ("WorksheetPages", "GraphPages", "MatrixPages"):
        pages: list[dict[str, Any]] = []
        for page in _required_collection_items(_required_attr(app, collection_name)):
            page_name = str(_required_attr(page, "Name", ""))
            page_layers: list[dict[str, Any]] = []
            for layer_position, raw_layer in enumerate(
                _required_collection_items(_required_attr(page, "Layers"))
            ):
                layer_name = str(_required_attr(raw_layer, "Name", ""))
                layer = raw_layer
                try:
                    if collection_name == "GraphPages":
                        layer = app.FindGraphLayer(f"[{page_name}]{layer_position + 1}") or raw_layer
                    elif collection_name == "WorksheetPages" and layer_name:
                        layer = app.FindWorksheet(f"[{page_name}]{layer_name}") or raw_layer
                    elif collection_name == "MatrixPages" and layer_name:
                        layer = app.FindMatrixSheet(f"[{page_name}]{layer_name}") or raw_layer
                except Exception:
                    layer = raw_layer
                layer_fingerprint: dict[str, Any] = {
                    "name": layer_name,
                    "long_name": str(_required_attr(layer, "LongName", "")),
                    "index": _required_attr(layer, "Index"),
                }
                if collection_name == "WorksheetPages":
                    columns = _required_attr(layer, "Columns") or _required_attr(
                        layer,
                        "DataObjectBases",
                    )
                    layer_fingerprint.update(
                        {
                            "rows": _required_attr(layer, "Rows"),
                            "cols": _required_attr(layer, "Cols"),
                            "columns": [
                                {
                                    "index": _required_attr(column, "Index", position),
                                    "name": str(_required_attr(column, "Name", "")),
                                    "long_name": str(_required_attr(column, "LongName", "")),
                                    "units": str(_required_attr(column, "Units", "")),
                                    "comments": str(_required_attr(column, "Comments", "")),
                                    "plot_designation": _required_attr(
                                        column,
                                        "PlotDesignation",
                                    ),
                                    "dataset_name": _origin_text(
                                        _safe_call(column, "GetDatasetName")
                                    ),
                                    "range": _origin_text(_required_attr(column, "Range")),
                                }
                                for position, column in enumerate(
                                    _required_collection_items(columns)
                                )
                            ],
                        }
                    )
                elif collection_name == "GraphPages":
                    plots = _required_collection_items(_required_attr(layer, "DataPlots"))
                    plot_fingerprints: list[dict[str, Any]] = []
                    for plot_position, plot in enumerate(plots, start=1):
                        base = {
                            "name": str(_required_attr(plot, "Name", "")),
                            "type": str(_required_attr(plot, "TypeName", "")),
                            "range": str(_required_attr(plot, "Range", "")),
                        }
                        described = _describe_data_plot(
                            app,
                            layer,
                            plot,
                            base,
                            graph_name=page_name,
                            layer_position=layer_position + 1,
                            plot_index=plot_position,
                            catalog=worksheet_catalog,
                        )
                        plot_fingerprints.append(
                            {
                                key: described.get(key)
                                for key in (
                                    "plot_index",
                                    "dataset_name",
                                    "source_workbook",
                                    "source_worksheet",
                                    "x_column",
                                    "x_long_name",
                                    "x_dataset_name",
                                    "x_range",
                                    "y_column",
                                    "y_long_name",
                                    "y_dataset_name",
                                    "y_range",
                                    "label_column",
                                    "label_long_name",
                                    "label_dataset_name",
                                    "label_range",
                                    "range",
                                )
                            }
                        )
                    layer_fingerprint["data_plots"] = plot_fingerprints
                else:
                    layer_fingerprint.update(
                        {
                            "rows": _required_attr(layer, "Rows"),
                            "cols": _required_attr(layer, "Cols"),
                        }
                    )
                page_layers.append(layer_fingerprint)
            pages.append(
                {
                    "name": page_name,
                    "long_name": str(_required_attr(page, "LongName", "")),
                    "layers": page_layers,
                }
            )
        fingerprint[collection_name] = sorted(
            pages,
            key=lambda item: (item["name"].casefold(), item["long_name"].casefold()),
        )
    return fingerprint


def _project_file_metadata(path: Path) -> dict[str, Any]:
    signature = project_signature(path)
    if signature is None:
        raise FileNotFoundError(f"Origin project does not exist: {path}")
    size_bytes, mtime_ns, digest = signature
    identity = _file_identity(path)
    return {
        "path": str(path),
        "size_bytes": size_bytes,
        "mtime_ns": mtime_ns,
        "sha256": digest,
        "file_identity": list(identity) if identity is not None else None,
    }


def _checked_column_index(sheet: Any, index: int, column: str | int) -> int:
    count = _safe_attr(sheet, "Cols")
    if index < 0 or (count is not None and index >= int(count)):
        raise LookupError(f"Origin worksheet column is out of range: {column}")
    return index


def _column_index(sheet: Any, column: str | int) -> int:
    if isinstance(column, int):
        return _checked_column_index(sheet, column, column)
    text = str(column).strip()
    if text.isdigit():
        return _checked_column_index(sheet, int(text), column)
    try:
        found = sheet.FindCol(text, 0, False)
        found_index = int(getattr(found, "Index", found)) if found is not None else -1
    except Exception:
        found_index = -1
    if found_index < 0:
        # Permit simple A/B/C references when labels are unavailable.
        if re.fullmatch(r"[A-Za-z]{1,3}", text):
            index = 0
            for char in text.upper():
                index = index * 26 + ord(char) - ord("A") + 1
            return _checked_column_index(sheet, index - 1, column)
        raise LookupError(f"Origin worksheet column not found: {column}")
    return _checked_column_index(sheet, found_index, column)


def _worksheet_column(sheet: Any, column: str | int) -> Any:
    column_index = _column_index(sheet, column)
    columns = _safe_attr(sheet, "Columns") or _safe_attr(sheet, "DataObjectBases")
    if columns is not None:
        try:
            resolved = columns.Item(column_index)
        except Exception:
            items = _collection_items(columns)
            resolved = items[column_index] if column_index < len(items) else None
        if resolved is not None:
            return resolved
    raise LookupError(f"Origin worksheet column was not found: {column}")


def _worksheet_reference(identifier: str) -> str:
    if identifier.startswith("worksheet_page:"):
        parts = identifier.split(":", 2)
        if len(parts) == 3:
            return f"[{parts[1]}]{parts[2]}"
    return identifier


def _resolve_worksheet(app: Any, identifier: str) -> Any:
    reference = _worksheet_reference(identifier)
    sheet = app.FindWorksheet(reference)
    if sheet is None:
        raise LookupError(f"Origin worksheet not found: {identifier}")
    return sheet


def _graph_reference(app: Any, identifier: str) -> str:
    if not identifier.startswith("graph_page:"):
        return identifier
    parts = identifier.split(":", 2)
    if len(parts) == 2:
        return parts[1]
    page_name, layer_name = parts[1], parts[2]
    match = re.search(r"(\d+)$", layer_name)
    if match:
        return f"[{page_name}]{int(match.group(1))}"
    try:
        page = app.GraphPages.Item(page_name)
        for position, layer in enumerate(_collection_items(_safe_attr(page, "Layers")), start=1):
            if str(_safe_attr(layer, "Name", "")) == layer_name:
                return f"[{page_name}]{position}"
    except Exception:
        pass
    return identifier


def _graph_page_name(identifier: str) -> str:
    if identifier.startswith("[") and "]" in identifier:
        return identifier[1:].split("]", 1)[0]
    if identifier.startswith("graph_page:"):
        return identifier.split(":", 2)[1]
    return identifier


def _resolve_graph_layer(app: Any, identifier: str) -> Any:
    reference = _graph_reference(app, identifier)
    layer = app.FindGraphLayer(reference)
    if layer is None:
        raise LookupError(f"Origin graph layer not found: {identifier}")
    return layer


def _graph_configuration_commands(options: Mapping[str, Any], labtalk: str | None) -> list[str]:
    unknown = sorted(set(options) - GRAPH_OPTION_KEYS)
    if unknown:
        raise ValueError(f"Unsupported graph options: {', '.join(unknown)}")
    commands: list[str] = []
    for axis in ("x", "y"):
        if f"{axis}_min" in options:
            commands.append(f"layer.{axis}.from={float(options[f'{axis}_min'])};")
        if f"{axis}_max" in options:
            commands.append(f"layer.{axis}.to={float(options[f'{axis}_max'])};")
        if f"{axis}_tick_step" in options:
            step = float(options[f"{axis}_tick_step"])
            if step <= 0:
                raise ValueError(f"{axis}_tick_step must be positive")
            commands.append(f"layer.{axis}.inc={step};")
        if f"{axis}_scale" in options:
            scale = str(options[f"{axis}_scale"]).lower()
            if scale not in SCALE_TYPES:
                raise ValueError(f"{axis}_scale must be linear, log10, ln, or log2")
            commands.append(f"layer.{axis}.type={SCALE_TYPES[scale]};")
        if f"{axis}_title" in options:
            flag = "xb" if axis == "x" else "yl"
            commands.append(f"label -{flag} {_labtalk_quote(str(options[f'{axis}_title']))};")
    if options.get("rescale"):
        commands.append("rescale;")
    if options.get("legend") is True:
        commands.append("legend;")
    elif options.get("legend") is False:
        commands.append("legend -d;")
    styles = options.get("plot_styles", [])
    if not isinstance(styles, list):
        raise ValueError("plot_styles must be a list")
    for style in styles:
        if not isinstance(style, Mapping):
            raise ValueError("Each plot_styles entry must be an object")
        unknown_style = sorted(set(style) - {"plot_index", "color_index", "line_connection"})
        if unknown_style:
            raise ValueError(f"Unsupported plot style options: {', '.join(unknown_style)}")
        plot_index = int(style.get("plot_index", 1))
        if plot_index < 1:
            raise ValueError("plot_index must be 1 or greater")
        dataset = f"%({plot_index},@D)"
        if "color_index" in style:
            color_index = int(style["color_index"])
            if color_index < 0:
                raise ValueError("color_index must be non-negative")
            commands.append(f"set {dataset} -c {color_index};")
        if "line_connection" in style:
            raw_connection = style["line_connection"]
            if isinstance(raw_connection, str):
                connection_name = raw_connection.lower()
                if connection_name not in LINE_CONNECTIONS:
                    raise ValueError(f"Unsupported line_connection: {raw_connection}")
                connection = LINE_CONNECTIONS[connection_name]
            else:
                connection = int(raw_connection)
            commands.append(f"set {dataset} -l {connection};")
    if labtalk:
        commands.append(labtalk)
    return commands


def _normalize_graph_data_binding(value: Any) -> dict[str, Any] | None:
    if value is None:
        return None
    if not isinstance(value, Mapping):
        raise ValueError("data_binding must be an object")
    unknown = sorted(
        set(value)
        - {"worksheet_name", "x_column", "y_columns", "label_column", "plot_type"}
    )
    if unknown:
        raise ValueError(f"Unsupported data_binding fields: {', '.join(unknown)}")
    worksheet_name = str(value.get("worksheet_name", "")).strip()
    y_columns = list(value.get("y_columns") or [])
    plot_type = str(value.get("plot_type", "line")).lower()
    if not worksheet_name or "x_column" not in value or not y_columns:
        raise ValueError(
            "data_binding requires worksheet_name, x_column, and at least one y_column"
        )
    if plot_type not in DATA_BINDING_PLOT_TYPES:
        raise ValueError("data_binding plot_type must be line, scatter, line_symbol, or bar")
    normalized = {
        "worksheet_name": worksheet_name,
        "x_column": value["x_column"],
        "y_columns": y_columns,
        "plot_type": plot_type,
    }
    if value.get("label_column") is not None:
        normalized["label_column"] = value["label_column"]
    return normalized


def _normalize_categorical_style(value: Any) -> dict[str, Any] | None:
    if value is None:
        return None
    if not isinstance(value, Mapping):
        raise ValueError("categorical_style must be an object")
    unknown = sorted(
        set(value) - {"plot_index", "worksheet_name", "category_column", "categories"}
    )
    if unknown:
        raise ValueError(f"Unsupported categorical_style fields: {', '.join(unknown)}")
    if "category_column" not in value:
        raise ValueError("categorical_style requires category_column")
    plot_index = int(value.get("plot_index", 1))
    if plot_index < 1:
        raise ValueError("categorical_style plot_index must be 1 or greater")
    worksheet_name = str(value.get("worksheet_name", "")).strip() or None
    raw_categories = value.get("categories")
    if not isinstance(raw_categories, Mapping) or not raw_categories:
        raise ValueError("categorical_style categories must be a non-empty object")
    categories: dict[str, dict[str, Any]] = {}
    for raw_category, raw_style in raw_categories.items():
        category = str(raw_category)
        if not category:
            raise ValueError("categorical_style category names must not be empty")
        if not isinstance(raw_style, Mapping):
            raise ValueError(f"Style for category {category!r} must be an object")
        unknown_style = sorted(set(raw_style) - {"color", "shape", "fill", "size"})
        if unknown_style:
            raise ValueError(
                f"Unsupported style fields for category {category!r}: {', '.join(unknown_style)}"
            )
        style: dict[str, Any] = {}
        if "color" in raw_style:
            color = str(raw_style["color"]).strip()
            if not re.fullmatch(r"#[0-9A-Fa-f]{6}", color):
                raise ValueError(f"Category {category!r} color must use #RRGGBB")
            style["color"] = color.upper()
        if "shape" in raw_style:
            shape = str(raw_style["shape"]).strip().lower()
            if shape not in CATEGORY_SHAPES:
                raise ValueError(
                    f"Category {category!r} shape must be one of: "
                    f"{', '.join(CATEGORY_SHAPES)}"
                )
            style["shape"] = shape
        if "fill" in raw_style:
            fill = str(raw_style["fill"]).strip().lower()
            if fill not in CATEGORY_FILLS:
                raise ValueError(
                    f"Category {category!r} fill must be one of: {', '.join(CATEGORY_FILLS)}"
                )
            style["fill"] = fill
        if "size" in raw_style:
            size = float(raw_style["size"])
            if size <= 0:
                raise ValueError(f"Category {category!r} size must be positive")
            style["size"] = size
        if not style:
            raise ValueError(f"Style for category {category!r} must set at least one property")
        categories[category] = style
    return {
        "plot_index": plot_index,
        "worksheet_name": worksheet_name,
        "category_column": value["category_column"],
        "categories": categories,
    }


def _normalize_categorical_legend(value: Any) -> dict[str, Any] | None:
    if value is None or isinstance(value, bool):
        return None
    if not isinstance(value, Mapping):
        raise ValueError("legend must be true, false, or an object")
    allowed = {
        "mode",
        "source",
        "plot_index",
        "position",
        "font_size",
        "line_spacing",
        "border",
        "background",
        "replace_existing",
        "show_all_categories",
    }
    unknown = sorted(set(value) - allowed)
    if unknown:
        raise ValueError(f"Unsupported categorical legend fields: {', '.join(unknown)}")
    mode = str(value.get("mode", "categorical")).lower()
    source = str(value.get("source", "plot_style_mapping")).lower()
    position = str(value.get("position", "top_right")).lower()
    background = str(value.get("background", "transparent")).lower()
    if mode != "categorical":
        raise ValueError("Structured legend mode must be categorical")
    if source != "plot_style_mapping":
        raise ValueError("Categorical legend source must be plot_style_mapping")
    if position != "top_right":
        raise ValueError("Only the verified categorical legend position top_right is supported")
    if background != "transparent":
        raise ValueError("Only the verified categorical legend background transparent is supported")
    plot_index = int(value.get("plot_index", 1))
    if plot_index < 1:
        raise ValueError("legend plot_index must be 1 or greater")
    font_size = value.get("font_size")
    if font_size is not None and float(font_size) <= 0:
        raise ValueError("legend font_size must be positive")
    line_spacing = value.get("line_spacing")
    if line_spacing is not None and float(line_spacing) <= 0:
        raise ValueError("legend line_spacing must be positive")
    border = value.get("border", False)
    replace_existing = value.get("replace_existing", True)
    show_all = value.get("show_all_categories", False)
    if not isinstance(border, bool):
        raise ValueError("legend border must be true or false")
    if not isinstance(replace_existing, bool):
        raise ValueError("legend replace_existing must be true or false")
    if not isinstance(show_all, bool):
        raise ValueError("legend show_all_categories must be true or false")
    return {
        "mode": mode,
        "source": source,
        "plot_index": plot_index,
        "position": position,
        "font_size": None if font_size is None else float(font_size),
        "line_spacing": None if line_spacing is None else float(line_spacing),
        "border": border,
        "background": background,
        "replace_existing": replace_existing,
        "show_all_categories": show_all,
    }


def _origin_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text == "###" or re.fullmatch(r"%\s*\([^)]*\)", text):
        return None
    return text


def _origin_data_reference(obj: Any) -> str | None:
    dataset_name = _origin_text(_safe_call(obj, "GetDatasetName"))
    if dataset_name:
        return dataset_name
    return _origin_text(_safe_attr(obj, "Range"))


def _data_plot_label_commands(plot: Any, label_column: Any) -> list[str]:
    plot_reference = _origin_data_reference(plot)
    label_reference = _origin_data_reference(label_column)
    if not plot_reference or not label_reference:
        raise LookupError("Origin did not expose stable dataset references for plot labels")
    return [
        f"set {plot_reference} -q 1;",
        f"set {plot_reference} -qm 5;",
        f"set {plot_reference} -j -qms %({label_reference}[i]$);",
    ]


def _parse_origin_column_reference(value: Any) -> tuple[str, str] | None:
    if not isinstance(value, str):
        return None
    match = re.fullmatch(
        r'\[([^\]]+)\]([^!]+)!Col\(\s*(?:"([^"]+)"|([^\)]+))\s*\)',
        value.strip(),
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    column = (match.group(3) or match.group(4) or "").strip()
    return f"[{match.group(1)}]{match.group(2)}", column


def _worksheet_column_catalog(app: Any) -> list[dict[str, Any]]:
    catalog: list[dict[str, Any]] = []
    for page in _required_collection_items(_required_attr(app, "WorksheetPages")):
        workbook_name = str(_safe_attr(page, "Name", ""))
        for raw_sheet in _required_collection_items(_required_attr(page, "Layers")):
            sheet_name = str(_safe_attr(raw_sheet, "Name", ""))
            sheet = raw_sheet
            if workbook_name and sheet_name:
                try:
                    sheet = app.FindWorksheet(f"[{workbook_name}]{sheet_name}") or raw_sheet
                except Exception:
                    sheet = raw_sheet
            columns = _required_attr(sheet, "Columns") or _required_attr(
                sheet, "DataObjectBases"
            )
            for position, column in enumerate(_required_collection_items(columns)):
                catalog.append(
                    {
                        "workbook": workbook_name,
                        "worksheet": sheet_name,
                        "sheet": sheet,
                        "column": column,
                        "index": _safe_attr(column, "Index", position),
                        "name": _origin_text(_safe_attr(column, "Name")),
                        "long_name": _origin_text(_safe_attr(column, "LongName")),
                        "dataset_name": _origin_text(_safe_call(column, "GetDatasetName")),
                        "range": _origin_text(_safe_attr(column, "Range")),
                        "plot_designation": _safe_attr(column, "PlotDesignation"),
                    }
                )
    return catalog


def _range_without_rows(value: str | None) -> str | None:
    if value is None:
        return None
    return re.sub(r"\[[^\[\]]*:[^\[\]]*\]$", "", value).strip()


def _column_source_matches(entry: Mapping[str, Any], *values: Any) -> bool:
    candidates = {
        str(candidate).casefold()
        for candidate in (
            entry.get("name"),
            entry.get("long_name"),
            entry.get("dataset_name"),
            _range_without_rows(entry.get("range")),
        )
        if candidate
    }
    for raw_value in values:
        value = _origin_text(raw_value)
        if not value:
            continue
        normalized = _range_without_rows(value)
        if normalized and normalized.casefold() in candidates:
            return True
        parsed = _parse_origin_column_reference(normalized)
        if parsed and parsed[1].casefold() in candidates:
            return True
    return False


def _catalog_column(
    catalog: Iterable[Mapping[str, Any]],
    *values: Any,
    workbook: str | None = None,
    worksheet: str | None = None,
) -> Mapping[str, Any] | None:
    for entry in catalog:
        if workbook and str(entry.get("workbook", "")).casefold() != workbook.casefold():
            continue
        if worksheet and str(entry.get("worksheet", "")).casefold() != worksheet.casefold():
            continue
        if _column_source_matches(entry, *values):
            return entry
    return None


def _categorical_style_commands(
    app: Any,
    layer: Any,
    style: Mapping[str, Any],
    data_binding: Mapping[str, Any] | None,
) -> tuple[list[str], dict[str, Any]]:
    plots = _collection_items(_safe_attr(layer, "DataPlots"))
    plot_index = int(style["plot_index"])
    if plot_index > len(plots):
        raise LookupError(f"Origin data plot {plot_index} does not exist in the graph layer")
    plot = plots[plot_index - 1]
    parsed_reference = _parse_origin_column_reference(style["category_column"])
    worksheet_name = style.get("worksheet_name")
    column_identifier: Any = style["category_column"]
    if parsed_reference:
        worksheet_name, column_identifier = parsed_reference
    elif not worksheet_name and data_binding:
        worksheet_name = data_binding["worksheet_name"]
    if not worksheet_name:
        catalog = _worksheet_column_catalog(app)
        source = _catalog_column(catalog, _origin_data_reference(plot))
        if source:
            worksheet_name = f"[{source['workbook']}]{source['worksheet']}"
    if not worksheet_name:
        raise ValueError(
            "categorical_style requires worksheet_name unless category_column is a full Origin range"
        )
    sheet = _resolve_worksheet(app, str(worksheet_name))
    column = _worksheet_column(sheet, column_identifier)
    category_reference = _origin_data_reference(column)
    plot_reference = None
    if data_binding:
        y_columns = data_binding["y_columns"]
        if plot_index > len(y_columns):
            raise LookupError(
                f"categorical_style plot_index {plot_index} has no matching data_binding Y column"
            )
        plot_sheet = _resolve_worksheet(app, data_binding["worksheet_name"])
        y_column = _worksheet_column(plot_sheet, y_columns[plot_index - 1])
        plot_reference = _origin_data_reference(y_column)
    if not plot_reference:
        plot_reference = _origin_data_reference(plot)
    if not category_reference or not plot_reference:
        raise LookupError("Origin did not expose stable dataset references for categorical styling")
    row_count = int(_safe_attr(sheet, "Rows", 0) or 0)
    category_values = [str(value) for value in _column_string_values(column, 0, row_count - 1)]
    categories = style["categories"]
    unmapped = sorted(set(category_values) - set(categories))
    if unmapped:
        raise ValueError(
            "categorical_style has no mapping for worksheet values: "
            + ", ".join(repr(value) for value in unmapped)
        )
    commands = [
        f"set {category_reference} -dc 1;",
        f"set {plot_reference} -ksn {category_reference};",
        f"set {plot_reference} -c {category_reference};",
        f"set {plot_reference} -csfn {category_reference};",
    ]
    category_rows: dict[str, int] = {}
    for row, category in enumerate(category_values, start=1):
        category_rows.setdefault(category, row)
        point_style = categories[category]
        if "shape" in point_style:
            commands.append(
                f"set {plot_reference} {row} -k {CATEGORY_SHAPES[point_style['shape']]};"
            )
        if "fill" in point_style:
            commands.append(
                f"set {plot_reference} {row} -kf {CATEGORY_FILLS[point_style['fill']]};"
            )
        if "color" in point_style:
            color = point_style["color"]
            commands.append(f'set {plot_reference} {row} -c color("{color}");')
            commands.append(f'set {plot_reference} {row} -csf color("{color}");')
        if "size" in point_style:
            commands.append(f"set {plot_reference} {row} -z {float(point_style['size'])};")
    counts = Counter(category_values)
    return commands, {
        "plot_index": plot_index,
        "worksheet_name": str(worksheet_name),
        "category_column": style["category_column"],
        "category_dataset": category_reference,
        "plot_dataset": plot_reference,
        "category_counts": dict(counts),
        "category_order": list(categories),
        "category_rows": category_rows,
        "unused_categories": [name for name in categories if name not in counts],
        "rows_styled": len(category_values),
    }


def _categorical_legend_commands(
    legend: Mapping[str, Any],
    categorical_result: Mapping[str, Any],
) -> list[str]:
    if not categorical_result["category_counts"]:
        raise LabTalkExecutionError("Categorical Legend has no populated category entries")
    if legend["show_all_categories"] and categorical_result["unused_categories"]:
        unused = ", ".join(categorical_result["unused_categories"])
        raise LabTalkExecutionError(
            f"Cannot render unused categorical Legend entries without sample points: {unused}"
        )
    show_all = 1 if legend["show_all_categories"] else 0
    commands = [f"legendcat mode:=2 combine:=1 showall:={show_all};"]
    if legend["font_size"] is not None:
        commands.append(f"Legend.fsize={legend['font_size']};")
    if legend["line_spacing"] is not None:
        commands.append(f"Legend.vgap={legend['line_spacing']};")
    commands.append(f"Legend.background={1 if legend['border'] else 0};")
    return commands


def _verify_categorical_legend(
    app: Any,
    layer: Any,
    replace_existing: bool,
    expected_entries: int | None = None,
) -> None:
    def verify_text(text: str | None) -> bool:
        if not text:
            return False
        normalized = re.sub(r"%\(CRLF\)", "\n", text, flags=re.IGNORECASE)
        entry_count = len([line for line in normalized.splitlines() if line.strip()])
        if expected_entries is None or entry_count == expected_entries:
            return True
        raise LabTalkExecutionError(
            "Origin categorical Legend entry count differs from the configured categories: "
            f"expected {expected_entries}, got {entry_count}"
        )

    graph_objects = _safe_attr(layer, "GraphObjects")
    items = _collection_items(graph_objects) if graph_objects is not None else []
    names = [str(_safe_attr(item, "Name", "")) for item in items]
    if replace_existing and any(name.casefold() == "categorykey" for name in names):
        raise LabTalkExecutionError("Origin did not remove the legacy CategoryKey object")
    enumerated_legends = [
        item
        for item, name in zip(items, names, strict=True)
        if name.casefold() == "legend"
    ]
    if len(enumerated_legends) > 1:
        raise LabTalkExecutionError("Origin left more than one named categorical Legend")

    legend = None
    if graph_objects is not None:
        try:
            item_method = getattr(graph_objects, "Item", None)
            if callable(item_method):
                legend = item_method("Legend")
            elif callable(graph_objects):
                legend = graph_objects("Legend")
        except Exception:
            legend = None
    if legend is None and enumerated_legends:
        legend = enumerated_legends[0]
    if legend is not None:
        legend_text = _origin_text(_safe_attr(legend, "Text"))
        if verify_text(legend_text):
            return

    executor = getattr(layer, "Execute", None)
    if not callable(executor):
        executor = getattr(app, "Execute", None)
    reader = getattr(layer, "LTStr", None)
    if not callable(reader):
        reader = getattr(app, "LTStr", None)
    if callable(executor) and callable(reader):
        prefix = f"cxlegend{uuid4().hex[:6]}"
        name_variable = f"{prefix}name"
        text_variable = f"{prefix}text"
        script = (
            f'{name_variable}$="%(Legend.name$)";'
            f'{text_variable}$="%(Legend.text$)";'
        )
        if executor(script) is not False:
            try:
                _origin_text(reader(name_variable))
                legend_text = _origin_text(reader(text_variable))
            finally:
                try:
                    executor(f"del -v {name_variable}$;del -v {text_variable}$;")
                except Exception:
                    pass
            if verify_text(legend_text):
                return
    raise LabTalkExecutionError("Origin did not expose a populated native categorical Legend")


def _execute_graph_commands(layer: Any, commands: list[str], failure_message: str) -> None:
    if not commands:
        return
    if layer.Execute("".join(commands)) is False:
        raise LabTalkExecutionError(failure_message)


def _plot_substitution_values(app: Any, layer: Any, plot_index: int) -> dict[str, str | None]:
    tokens = {
        "source_workbook": f"%({plot_index},@W)",
        "source_worksheet": f"%({plot_index},@WS)",
        "x_dataset_name": f"%({plot_index}X,@D)",
        "x_range": f"%({plot_index}X,@R)",
        "x_long_name": f"%({plot_index}X,@L)",
        "y_dataset_name": f"%({plot_index}Y,@D)",
        "y_range": f"%({plot_index}Y,@R)",
        "y_long_name": f"%({plot_index}Y,@L)",
        "label_dataset_name": f"%({plot_index}L,@D)",
        "label_range": f"%({plot_index}L,@R)",
        "label_long_name": f"%({plot_index}L,@L)",
    }
    prefix = f"cx{uuid4().hex[:6]}"
    variables = {field: f"{prefix}{position}" for position, field in enumerate(tokens)}
    script = "".join(
        f'{variables[field]}$="{token}";' for field, token in tokens.items()
    )
    executor = getattr(layer, "Execute", None)
    if not callable(executor):
        return {field: None for field in tokens}
    try:
        if executor(script) is False:
            return {field: None for field in tokens}
        reader = getattr(layer, "LTStr", None)
        if not callable(reader):
            reader = getattr(app, "LTStr", None)
        if not callable(reader):
            return {field: None for field in tokens}
        values = {
            field: _origin_text(reader(variable)) for field, variable in variables.items()
        }
        try:
            executor("".join(f"del -v {variable}$;" for variable in variables.values()))
        except Exception:
            pass
        return values
    except Exception:
        return {field: None for field in tokens}


def _custom_label_dataset_reference(value: Any) -> str | None:
    if value is None:
        return None
    match = re.fullmatch(r"%\((.+)\[i\]\$\)", str(value).strip(), flags=re.IGNORECASE)
    return _origin_text(match.group(1)) if match else None


def _plot_custom_label_dataset_reference(app: Any, layer: Any, plot_reference: str) -> str | None:
    executor = getattr(layer, "Execute", None)
    if not callable(executor):
        return None
    reader = getattr(layer, "LTStr", None)
    if not callable(reader):
        reader = getattr(app, "LTStr", None)
    if not callable(reader):
        return None
    variable = f"cxlabel{uuid4().hex[:6]}"
    try:
        if executor(f"get {plot_reference} -qms {variable}$;") is False:
            return None
        return _custom_label_dataset_reference(reader(variable))
    except Exception:
        return None
    finally:
        try:
            executor(f"del -v {variable}$;")
        except Exception:
            pass


def _designated_column(
    catalog: Iterable[Mapping[str, Any]],
    designation: int,
    *,
    workbook: str | None,
    worksheet: str | None,
) -> Mapping[str, Any] | None:
    for entry in catalog:
        if workbook and str(entry.get("workbook", "")).casefold() != workbook.casefold():
            continue
        if worksheet and str(entry.get("worksheet", "")).casefold() != worksheet.casefold():
            continue
        try:
            if int(entry.get("plot_designation")) == designation:
                return entry
        except (TypeError, ValueError):
            continue
    return None


def _role_column_name(entry: Mapping[str, Any] | None, range_value: str | None) -> str | None:
    if entry and entry.get("name"):
        return str(entry["name"])
    value = _origin_text(range_value)
    if value and re.fullmatch(r"[A-Za-z]{1,3}", value):
        return value.upper()
    parsed = _parse_origin_column_reference(value)
    return parsed[1] if parsed else None


def _describe_data_plot(
    app: Any,
    layer: Any,
    source: Any,
    base: Mapping[str, Any],
    *,
    graph_name: str,
    layer_position: int,
    plot_index: int,
    catalog: list[dict[str, Any]],
) -> dict[str, Any]:
    substitutions = _plot_substitution_values(app, layer, plot_index)
    dataset_name = _origin_text(_safe_call(source, "GetDatasetName")) or _origin_text(
        base.get("name")
    )
    workbook = substitutions["source_workbook"]
    worksheet = substitutions["source_worksheet"]
    y_entry = _catalog_column(
        catalog,
        substitutions["y_dataset_name"],
        substitutions["y_range"],
        dataset_name,
        workbook=workbook,
        worksheet=worksheet,
    )
    if y_entry is None:
        y_entry = _catalog_column(
            catalog,
            substitutions["y_dataset_name"],
            substitutions["y_range"],
            dataset_name,
        )
    if y_entry:
        workbook = str(y_entry["workbook"]) or workbook
        worksheet = str(y_entry["worksheet"]) or worksheet
    x_entry = _catalog_column(
        catalog,
        substitutions["x_dataset_name"],
        substitutions["x_range"],
        workbook=workbook,
        worksheet=worksheet,
    )
    if x_entry is None:
        x_entry = _designated_column(
            catalog,
            ORIGIN_PLOT_DESIGNATION_X,
            workbook=workbook,
            worksheet=worksheet,
        )
    label_form = _origin_text(_safe_call(source, "GetStrProp", "label.form"))
    label_visible = _safe_call(source, "GetNumProp", "label.show")
    label_reference = _custom_label_dataset_reference(label_form)
    if label_visible and not label_reference and dataset_name:
        label_reference = _plot_custom_label_dataset_reference(
            app,
            layer,
            dataset_name,
        )
    has_label_role = any(
        substitutions[field]
        for field in ("label_dataset_name", "label_range", "label_long_name")
    ) or bool(label_visible and (label_reference or label_form))
    label_entry = None
    if has_label_role:
        label_entry = _catalog_column(
            catalog,
            substitutions["label_dataset_name"],
            substitutions["label_range"],
            label_reference,
            label_form,
            workbook=workbook,
            worksheet=worksheet,
        )
    label_source_status = "resolved" if label_entry is not None else (
        "unresolved" if has_label_role else "none"
    )
    result = dict(base)
    result.update(
        {
            "type": "DataPlot",
            "origin_type_name": _origin_text(base.get("type")),
            "graph": graph_name,
            "layer": layer_position,
            "plot_index": plot_index,
            "dataset_name": dataset_name,
            "source_workbook": workbook,
            "source_worksheet": worksheet,
            "x_column": _role_column_name(x_entry, substitutions["x_range"]),
            "x_long_name": substitutions["x_long_name"]
            or (str(x_entry["long_name"]) if x_entry and x_entry.get("long_name") else None),
            "x_dataset_name": substitutions["x_dataset_name"]
            or (str(x_entry["dataset_name"]) if x_entry and x_entry.get("dataset_name") else None),
            "x_range": substitutions["x_range"]
            or (str(x_entry["range"]) if x_entry and x_entry.get("range") else None),
            "y_column": _role_column_name(y_entry, substitutions["y_range"]),
            "y_long_name": substitutions["y_long_name"]
            or (str(y_entry["long_name"]) if y_entry and y_entry.get("long_name") else None),
            "y_dataset_name": substitutions["y_dataset_name"]
            or (str(y_entry["dataset_name"]) if y_entry and y_entry.get("dataset_name") else dataset_name),
            "y_range": substitutions["y_range"]
            or (str(y_entry["range"]) if y_entry and y_entry.get("range") else None),
            "label_source_status": label_source_status,
            "label_column": _role_column_name(label_entry, substitutions["label_range"])
            if label_entry is not None
            else None,
            "label_long_name": (
                substitutions["label_long_name"]
                or (
                    str(label_entry["long_name"])
                    if label_entry and label_entry.get("long_name")
                    else None
                )
            )
            if label_entry is not None
            else None,
            "label_dataset_name": (
                substitutions["label_dataset_name"]
                or (
                    str(label_entry["dataset_name"])
                    if label_entry and label_entry.get("dataset_name")
                    else None
                )
            )
            if label_entry is not None
            else None,
            "label_range": (
                substitutions["label_range"]
                or (str(label_entry["range"]) if label_entry and label_entry.get("range") else None)
            )
            if label_entry is not None
            else None,
        }
    )
    return result


def _normalize_analysis_filters(
    filters: list[Mapping[str, Any]] | None,
) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for condition in filters or []:
        unknown = sorted(set(condition) - {"column", "operator", "value"})
        if unknown:
            raise ValueError(f"Unsupported filter fields: {', '.join(unknown)}")
        column = str(condition.get("column", "")).lower()
        operator = str(condition.get("operator", "")).lower()
        if column not in {"x", "y"}:
            raise ValueError("Filter column must be x or y")
        if operator not in FILTER_OPERATORS:
            raise ValueError("Filter operator must be gt, ge, lt, le, eq, or ne")
        normalized.append(
            {"column": column, "operator": operator, "value": float(condition["value"])}
        )
    return normalized


def _serialized_operation(method: Callable[..., ResultEnvelope]) -> Callable[..., ResultEnvelope]:
    signature = inspect.signature(method)

    @wraps(method)
    def locked(self: "OriginController", *args: Any, **kwargs: Any) -> ResultEnvelope:
        task_id = uuid4().hex
        started = time.perf_counter()
        with self._operation_lock:
            result = method(self, *args, **kwargs)
        if result.origin_version is None and self._origin_version:
            result = replace(result, origin_version=self._origin_version)
        try:
            bound = signature.bind_partial(self, *args, **kwargs)
            object_name = next(
                (
                    str(bound.arguments[key])
                    for key in ("name", "worksheet_name", "graph_name")
                    if key in bound.arguments
                ),
                None,
            )
        except Exception:
            object_name = None
        logger.info(
            json.dumps(
                {
                    "event": "origin_operation",
                    "task_id": task_id,
                    "operation": method.__name__,
                    "success": result.success,
                    "duration_ms": int((time.perf_counter() - started) * 1000),
                    "origin_version": result.origin_version,
                    "project_path": str(self._project_path) if self._project_path else None,
                    "object_name": object_name,
                    "warnings": result.warnings,
                    "error_code": result.error_code,
                    "error_message": result.error_message,
                },
                ensure_ascii=True,
            )
        )
        return result

    return locked


def _coerce_delimited_cell(value: str) -> Any:
    text = value.strip()
    if text == "":
        return None
    try:
        return int(text)
    except ValueError:
        try:
            return float(text)
        except ValueError:
            return value


def _split_header(
    rows: list[list[Any]],
    has_header: bool | None,
) -> tuple[list[str] | None, list[list[Any]]]:
    detected = has_header
    if detected is None:
        detected = False
        if len(rows) >= 2:
            first = rows[0]
            second = rows[1]
            first_has_text = any(isinstance(value, str) and value.strip() for value in first)
            second_values = [value for value in second if value is not None]
            second_is_numeric = bool(second_values) and all(
                isinstance(value, (int, float)) and not isinstance(value, bool)
                for value in second_values
            )
            detected = first_has_text and second_is_numeric
    if not detected or not rows:
        return None, rows
    return ["" if value is None else str(value) for value in rows[0]], rows[1:]


def _exception_error_code(exc: Exception) -> str:
    hresult = getattr(exc, "hresult", None)
    if hresult is None and exc.args and isinstance(exc.args[0], int):
        hresult = exc.args[0]
    if isinstance(hresult, int):
        if hresult > 0x7FFFFFFF:
            hresult -= 0x100000000
        if hresult in HRESULT_ERROR_CODES:
            return HRESULT_ERROR_CODES[hresult]
    if isinstance(exc, OSError) and getattr(exc, "winerror", None) in {32, 33}:
        return "FILE_LOCKED"
    if isinstance(exc, FileNotFoundError):
        return "PATH_NOT_FOUND"
    if isinstance(exc, PermissionError):
        return "PERMISSION_DENIED"
    return "COM_ERROR"


class OriginController:
    """Own the COM proxy only inside a serialized worker thread."""

    def __init__(
        self,
        *,
        worker: Any | None = None,
        dispatch_factory: Callable[[str, bool], Any] | None = None,
        process_snapshot: Callable[[], set[int]] | None = None,
        process_terminator: Callable[[int], None] | None = None,
        config: PluginConfig | None = None,
        sleep_fn: Callable[[float], None] | None = None,
    ) -> None:
        self.worker = worker or SerialComWorker()
        self.dispatch_factory = dispatch_factory or _default_dispatch
        self.process_snapshot = process_snapshot or _default_process_snapshot
        # Kept as an injection point for compatibility with earlier callers.
        # PID-only termination is intentionally never used because it cannot
        # prove that a process belongs to this COM activation.
        self.process_terminator = process_terminator
        self.sleep_fn = sleep_fn or time.sleep
        self.config = config or PluginConfig.from_environment(Path(__file__).resolve().parents[3])
        self._operation_lock = RLock()
        self.sessions = SessionManager()
        self._app: Any | None = None
        self._session_id: str | None = None
        self._project_path: Path | None = None
        self._source_path: Path | None = None
        self._protected_source_paths: set[Path] = set()
        self._protected_source_identities: set[tuple[int, int]] = set()
        self._origin_version: str | None = None
        self._exclusive = False
        self._poisoned = False
        self._retired = False
        self._current_stage: str | None = None
        self._analysis_operations = AnalysisOperationRegistry()

    @property
    def session_id(self) -> str | None:
        return self._session_id

    def _ensure_worker(self) -> None:
        start = getattr(self.worker, "start", None)
        if callable(start) and not getattr(self.worker, "is_running", False):
            start()

    def _submit(
        self,
        fn: Callable[[], Any],
        *,
        timeout: float | None = None,
        allow_poisoned: bool = False,
        retryable: bool = False,
        stage: str | None = None,
    ) -> ResultEnvelope:
        started = time.perf_counter()
        if stage is not None:
            self._current_stage = stage

        def context_data(data: Any = None) -> Any:
            context: dict[str, Any] = {"stage": self._current_stage}
            if self._session_id:
                context["session_id"] = self._session_id
                try:
                    record = self.sessions.get(self._session_id)
                    context.update(
                        {
                            "owned": record.owned,
                            "owned_pid": record.pid if record.owned else None,
                        }
                    )
                except Exception:
                    pass
            if isinstance(data, dict):
                return {**context, **data}
            return data if data is not None else context

        if self._retired:
            return ResultEnvelope.fail(
                "CONTROLLER_RETIRED",
                "This controller was abandoned after a poisoned COM session",
                duration_ms=0,
            )
        if self._poisoned and not allow_poisoned:
            return ResultEnvelope.fail(
                "SESSION_POISONED",
                "A previous COM call timed out; restart the session before issuing more operations",
                duration_ms=0,
            )
        warnings: list[str] = []
        for attempt in range(2):
            try:
                self._ensure_worker()
                data = self.worker.submit(
                    fn,
                    timeout=self.config.operation_timeout_s if timeout is None else timeout,
                )
                duration_ms = int((time.perf_counter() - started) * 1000)
                if isinstance(data, ResultEnvelope):
                    return replace(
                        data,
                        data=context_data(data.data),
                        warnings=[*warnings, *data.warnings],
                        duration_ms=data.duration_ms if data.duration_ms is not None else duration_ms,
                    )
                return ResultEnvelope.ok(
                    context_data(data), warnings=warnings, duration_ms=duration_ms
                )
            except OriginTimeoutError as exc:
                self._poisoned = True
                return ResultEnvelope.fail(
                    exc.code,
                    exc.message,
                    data=context_data(),
                    warnings=warnings,
                    duration_ms=int((time.perf_counter() - started) * 1000),
                )
            except OriginAutomationError as exc:
                return ResultEnvelope.fail(
                    exc.code,
                    exc.message,
                    data=context_data(),
                    warnings=warnings,
                    duration_ms=int((time.perf_counter() - started) * 1000),
                )
            except Exception as exc:
                error_code = _exception_error_code(exc)
                if error_code in {
                    "RPC_DISCONNECTED",
                    "RPC_SERVER_UNAVAILABLE",
                    "CO_E_SERVER_EXEC_FAILURE",
                }:
                    termination_warning = self._owned_process_termination_warning()
                    if termination_warning:
                        warnings.append(termination_warning)
                        error_code = "ORIGIN_PROCESS_TERMINATED"
                        self._poisoned = True
                    elif error_code == "RPC_DISCONNECTED":
                        self._poisoned = True
                if retryable and attempt == 0 and error_code in RETRYABLE_READ_ERRORS:
                    warnings.append(f"Retried once after {error_code}")
                    self.sleep_fn(0.25)
                    continue
                return ResultEnvelope.fail(
                    error_code,
                    str(exc),
                    data=context_data(),
                    warnings=warnings,
                    duration_ms=int((time.perf_counter() - started) * 1000),
                )
        raise AssertionError("unreachable")

    def _guard_mutation(self) -> ResultEnvelope | None:
        if not self._session_id:
            return ResultEnvelope.fail(
                "NO_ACTIVE_SESSION",
                "Start an owned Origin session before issuing a mutating operation",
            )
        record = self.sessions.get(self._session_id)
        if not record.owned:
            return ResultEnvelope.fail(
                "ATTACHED_SESSION_PROTECTED",
                "Mutating Origin operations are disabled for an attached user-owned session",
            )
        return None

    def _require_app(self) -> Any:
        if self._app is None or not self._session_id:
            raise NoActiveSessionError("Start or attach to an Origin session first")
        return self._app

    def _clear_session_state(self) -> None:
        self._app = None
        self._session_id = None
        self._project_path = None
        self._source_path = None
        self._protected_source_paths.clear()
        self._protected_source_identities.clear()
        self._exclusive = False
        self._poisoned = False
        self._analysis_operations = AnalysisOperationRegistry()

    def _owned_process_termination_warning(self) -> str | None:
        if not self._session_id:
            return None
        try:
            record = self.sessions.get(self._session_id)
            current_pids = set(self.process_snapshot())
        except Exception:
            return None
        if not record.owned or record.pid is None or record.pid in current_pids:
            return None
        try:
            from ..tools.health import find_origin_scheduled_tasks

            tasks = find_origin_scheduled_tasks()
        except Exception:
            tasks = []
        task_note = f" Candidate cleanup tasks: {', '.join(tasks)}." if tasks else ""
        return f"Plugin-owned Origin PID {record.pid} disappeared during the COM call.{task_note}"

    def _wait_for_pid_exit(
        self,
        pid: int,
        *,
        attempts: int = 100,
        interval_s: float = 0.2,
    ) -> tuple[bool, set[int] | None]:
        last_snapshot: set[int] | None = None
        for attempt in range(attempts):
            try:
                last_snapshot = set(self.process_snapshot())
            except Exception:
                return False, last_snapshot
            if pid not in last_snapshot:
                return True, last_snapshot
            if attempt + 1 < attempts:
                self.sleep_fn(interval_s)
        return False, last_snapshot

    @_serialized_operation
    def start(
        self,
        *,
        progid: str | None = None,
        visible: bool | None = None,
        attach: bool | None = None,
        exclusive: bool = False,
    ) -> ResultEnvelope:
        if self._session_id and self._app is not None:
            record = self.sessions.get(self._session_id)
            requested_attach = self.config.attach if attach is None else attach
            requested_visible = self.config.visible if visible is None else visible
            conflicts = []
            if progid is not None and progid != record.progid:
                conflicts.append("progid")
            if bool(requested_attach) == record.owned:
                conflicts.append("attach")
            if bool(requested_visible) != record.visible:
                conflicts.append("visible")
            if bool(exclusive) != self._exclusive:
                conflicts.append("exclusive")
            if conflicts:
                return ResultEnvelope.fail(
                    "SESSION_ALREADY_ACTIVE_CONFLICT",
                    f"An Origin session is already active with conflicting options: {', '.join(conflicts)}",
                )
            return ResultEnvelope.ok(
                {
                    "session_id": self._session_id,
                    "already_active": True,
                    "progid": record.progid,
                    "pid": record.pid,
                    "owned": record.owned,
                    "visible": record.visible,
                    "exclusive_requested": exclusive,
                    "exclusive": self._exclusive,
                    "exclusive_mode": "si_session_lock" if self._exclusive else None,
                    "process_isolated": record.owned and record.pid is not None,
                }
            )
        visible = self.config.visible if visible is None else visible
        attach = self.config.attach if attach is None else attach
        progid = progid or self.config.origin_progid
        registrations = discover_registrations()
        if progid is None:
            if attach:
                progid = "Origin.ApplicationSI"
            elif exclusive:
                selected = next(
                    (
                        item
                        for item in registrations
                        if item.available and item.progid in SINGLE_INSTANCE_PROGIDS
                    ),
                    None,
                )
                progid = selected.progid if selected else "Origin.ApplicationSI"
            else:
                selected = next((item for item in registrations if item.available), None)
                progid = selected.progid if selected else OWNED_INSTANCE_PROGID
        if exclusive and not attach:
            return ResultEnvelope.fail(
                "EXCLUSIVE_REQUIRES_ATTACH",
                "exclusive=true is only supported for an explicitly attached read-only SI/COMSI session",
                data={"exclusive_requested": True, "attach_requested": False, "progid": progid},
            )
        if attach and progid not in SINGLE_INSTANCE_PROGIDS:
            return ResultEnvelope.fail(
                "ATTACH_UNSUPPORTED",
                "attach=true requires Origin.ApplicationSI or Origin.ApplicationCOMSI",
                data={"attach_requested": True, "progid": progid},
            )
        if not attach and progid != OWNED_INSTANCE_PROGID:
            return ResultEnvelope.fail(
                "OWNED_PROGID_UNSUPPORTED",
                "Owned sessions require Origin.Application with DispatchEx fresh-instance activation",
                data={"attach_requested": False, "progid": progid},
            )
        if exclusive and progid not in SINGLE_INSTANCE_PROGIDS:
            return ResultEnvelope.fail(
                "EXCLUSIVE_UNSUPPORTED",
                "exclusive=true requires Origin.ApplicationSI or Origin.ApplicationCOMSI",
                data={"exclusive_requested": True, "progid": progid},
            )
        selected_registration = next(
            (item for item in registrations if item.progid == progid and item.available),
            None,
        )
        version_hint = (
            executable_version(selected_registration.server_path) if selected_registration else None
        )

        def activate() -> dict[str, Any]:
            before = set(self.process_snapshot())
            if attach and not before:
                raise NoExistingOriginError(
                    "attach=true requires an already running Origin instance"
                )
            try:
                app = self.dispatch_factory(progid, attach)
            except Exception as activation_error:
                warnings: list[str] = []
                try:
                    after_failure = set(self.process_snapshot())
                except Exception as snapshot_error:
                    after_failure = None
                    warnings.append(
                        f"Could not audit Origin PIDs after failed COM activation: {snapshot_error}"
                    )
                new_after_failure = (
                    after_failure - before if after_failure is not None else set()
                )
                if new_after_failure:
                    warnings.append(
                        "New Origin PIDs appeared during failed activation and were not terminated "
                        "because COM returned no proxy and PID ownership is unproven"
                    )
                return ResultEnvelope.fail(
                    _exception_error_code(activation_error),
                    str(activation_error),
                    data={
                        "progid": progid,
                        "attach_requested": attach,
                        "pids_before": sorted(before),
                        "pids_after": (
                            sorted(after_failure) if after_failure is not None else None
                        ),
                        "new_pids": sorted(new_after_failure),
                        "activation_cleanup_attempted": False,
                        "activation_cleanup_confirmed": False,
                        "pids_after_cleanup": (
                            sorted(after_failure) if after_failure is not None else None
                        ),
                        "pid_binding_confirmed": False,
                    },
                    warnings=warnings,
                    origin_version=version_hint,
                )
            after = set(self.process_snapshot())
            new_pids = after - before
            # An SI attach can also create a new instance. Wait briefly so a
            # delayed process is classified correctly before granting rights.
            attempts = 49 if not attach else 9
            for _ in range(attempts):
                if new_pids:
                    break
                if not attach and len(new_pids) > 1:
                    break
                self.sleep_fn(0.1)
                try:
                    after = set(self.process_snapshot())
                except (StopIteration, RuntimeError):
                    break
                new_pids = after - before
            if attach and new_pids:
                raise OwnershipUnverifiedError(
                    "An Origin PID appeared during attach; refusing to claim or close the ambiguous proxy"
                )
            if len(new_pids) > 1:
                raise OwnershipUnverifiedError(
                    "COM activation created multiple Origin processes; ownership is ambiguous"
                )
            fresh_instance_activation = not attach and progid == OWNED_INSTANCE_PROGID
            owned = fresh_instance_activation and len(new_pids) == 1
            if not attach and not owned:
                raise OwnershipUnverifiedError(
                    "DispatchEx returned a fresh Origin proxy but exactly one new audit PID was not observed"
                )
            if owned:
                try:
                    app.Visible = 1 if visible else 0
                except Exception:
                    pass
            version = _safe_attr(app, "Version", _safe_attr(app, "VersionS"))
            if version not in {None, ""}:
                self._origin_version = str(version)
            elif version_hint:
                self._origin_version = version_hint
            if exclusive:
                try:
                    begin_session = getattr(app, "BeginSession", None)
                    if not callable(begin_session):
                        raise OwnershipUnverifiedError(
                            f"{progid} does not expose BeginSession; exclusive lock was not acquired"
                        )
                    lock_result = begin_session()
                    if lock_result is False or lock_result == 0:
                        raise OwnershipUnverifiedError(
                            "Origin rejected the exclusive BeginSession lock"
                        )
                except Exception:
                    raise
                self._exclusive = True
            self._app = app
            self._poisoned = False
            self._retired = False
            self._protected_source_paths.clear()
            self._protected_source_identities.clear()
            pid = next(iter(new_pids)) if owned else (next(iter(before)) if len(before) == 1 else None)
            record = self.sessions.register(
                progid=progid,
                pid=pid,
                owned=owned,
                visible=visible,
            )
            self._session_id = record.session_id
            return {
                "session_id": record.session_id,
                "progid": progid,
                "pid": pid,
                "owned": owned,
                "visible": visible,
                "attach_requested": attach,
                "exclusive_requested": exclusive,
                "exclusive": self._exclusive,
                "exclusive_mode": "si_session_lock" if self._exclusive else None,
                "process_isolated": owned and pid is not None,
                "activation_mode": (
                    "dispatch_ex_fresh_instance" if owned else "si_read_only_attach"
                ),
                "ownership_basis": (
                    "Origin.Application DispatchEx fresh-instance contract"
                    if owned
                    else "Explicit SI/COMSI attachment"
                ),
                "pids_before": sorted(before),
                "pids_after": sorted(after),
                "new_pids": sorted(new_pids),
                "observed_new_pid": next(iter(new_pids)) if len(new_pids) == 1 else None,
                "pid_binding_confirmed": False,
            }

        return self._submit(activate)

    @_serialized_operation
    def abandon_poisoned_session(self) -> ResultEnvelope:
        """Quarantine a timed-out worker without calling into its blocked STA thread."""

        if not self._poisoned:
            return ResultEnvelope.fail(
                "RECOVERY_NOT_REQUIRED",
                "The active controller is not poisoned by a COM timeout",
            )
        session_id = self._session_id
        record = self.sessions.get(session_id) if session_id else None
        try:
            current_pids = set(self.process_snapshot())
        except Exception:
            current_pids = None
        owned_pid = record.pid if record and record.owned else None
        process_still_running = (
            owned_pid in current_pids
            if owned_pid is not None and current_pids is not None
            else None
        )
        if record is not None:
            record.active = False

        # Do not stop or reuse this worker. It may still be blocked inside COM.
        # A server-level recovery replaces the complete controller so any late
        # completion can mutate only this retired instance.
        self._app = None
        self._session_id = None
        self._project_path = None
        self._source_path = None
        self._protected_source_paths.clear()
        self._protected_source_identities.clear()
        self._exclusive = False
        self._retired = True
        return ResultEnvelope.ok(
            {
                "recovered": True,
                "previous_session_id": session_id,
                "owned": bool(record and record.owned),
                "owned_pid": owned_pid,
                "worker_abandoned": True,
                "process_still_running": process_still_running,
                "process_cleanup_attempted": False,
                "process_cleanup_confirmed": process_still_running is False,
                "cleanup_policy": "no_pid_termination_without_confirmed_proxy_binding",
                "stage": "controller_replaced",
            },
            warnings=(
                [
                    f"Owned Origin PID {owned_pid} is still present; it was not terminated "
                    "because proxy-to-PID binding is not independently confirmed"
                ]
                if process_still_running
                else []
            ),
        )

    @_serialized_operation
    def open_project(
        self,
        *,
        source_path: str,
        working_copy_path: str | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        source = Path(source_path).expanduser().resolve()
        if working_copy_path is None:
            working = source.with_name(f"{source.stem}.codex-working-{uuid4().hex[:8]}{source.suffix}")
        else:
            working = Path(working_copy_path).expanduser().resolve()
        try:
            self._current_stage = "copy"
            copied = copy_project(source, working)
        except (OSError, ValueError) as exc:
            return ResultEnvelope.fail("PROJECT_COPY_FAILED", str(exc))
        # A COM Load can partially switch projects before raising. Protect the
        # original as soon as the working copy exists, for the whole session.
        self._protected_source_paths.add(source)
        source_identity = _file_identity(source)
        if source_identity is not None:
            self._protected_source_identities.add(source_identity)

        def load() -> dict[str, Any]:
            app = self._require_app()
            # Load can switch the active project before reporting an error.
            # Until it succeeds, no prior source remains authorized for replacement.
            self._project_path = None
            self._source_path = None
            self._current_stage = "load"
            loaded = app.Load(str(copied), False)
            if loaded is False:
                raise OSError(f"Origin rejected project load: {copied}")
            self._current_stage = "activate"
            self._project_path = copied
            self._source_path = source
            self._current_stage = "validate"
            return {
                "source_path": str(source),
                "working_copy_path": str(copied),
                "completed_stages": ["copy", "load", "activate", "validate"],
                "working_copy_exists": copied.is_file(),
            }

        result = self._submit(load, stage="load")
        if result.success:
            result = ResultEnvelope.ok(
                result.data,
                warnings=["Source project was copied before COM opened it"],
                artifacts=[Artifact(str(copied), "origin_project")],
                duration_ms=result.duration_ms,
            )
        return result

    @_serialized_operation
    def save_project_copy(self, *, target_path: str, overwrite: bool = False) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        target = Path(target_path).expanduser().resolve()
        target_identity = _file_identity(target)
        if any(
            _same_file_path(target, source) for source in self._protected_source_paths
        ) or (
            target_identity is not None
            and target_identity in self._protected_source_identities
        ):
            return ResultEnvelope.fail(
                "SOURCE_OVERWRITE_BLOCKED",
                "Save target resolves to the original source project",
            )
        if self._project_path and target == self._project_path:
            return ResultEnvelope.fail("SOURCE_OVERWRITE_BLOCKED", "Save target must differ from the working project")
        if target.exists() and not overwrite:
            return ResultEnvelope.fail("OUTPUT_EXISTS", f"Refusing to overwrite existing output: {target}")
        before_signature = project_signature(target)

        def save() -> dict[str, Any]:
            app = self._require_app()
            target.parent.mkdir(parents=True, exist_ok=True)
            run = getattr(app, "Run", None)
            if callable(run) and run() is False:
                raise OSError("Origin did not finish pending recalculation before save")
            saved = app.Save(str(target))
            if saved not in (True, 1):
                raise ProjectSaveUnconfirmedError(
                    f"Origin did not explicitly confirm the project save: {target}"
                )
            if not validate_project_artifact(target):
                raise ProjectSaveUnconfirmedError(
                    f"Origin did not create a plausibly sized project file: {target}"
                )
            after_signature = project_signature(target)
            if before_signature is not None and after_signature == before_signature:
                raise ProjectSaveUnconfirmedError(
                    f"Origin reported success but the existing project file did not change: {target}"
                )
            modified = _safe_attr(app, "IsModified")
            if modified is True:
                raise ProjectSaveUnconfirmedError(
                    "Origin still reports unsaved project changes after Save"
                )
            self._project_path = target
            return {
                "path": str(target),
                "size_bytes": target.stat().st_size,
                "modified": modified,
                "project_name": str(_safe_attr(app, "Name", "")),
            }

        result = self._submit(save)
        if result.success:
            result = ResultEnvelope.ok(
                result.data,
                artifacts=[Artifact(str(target), "origin_project")],
                duration_ms=result.duration_ms,
            )
        return result

    @_serialized_operation
    def save_and_replace_source(
        self,
        *,
        source_path: str,
        overwrite: bool = False,
        allow_source_overwrite: bool = False,
        expected_source_sha256: str,
        keep_backup: bool = True,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        if not overwrite or not allow_source_overwrite:
            return ResultEnvelope.fail(
                "SOURCE_REPLACE_CONFIRMATION_REQUIRED",
                "Source replacement requires overwrite=true and allow_source_overwrite=true",
            )
        expected_digest = str(expected_source_sha256).strip().lower()
        if not re.fullmatch(r"[0-9a-f]{64}", expected_digest):
            return ResultEnvelope.fail(
                "INVALID_SOURCE_SHA256",
                "expected_source_sha256 must be a 64-character SHA-256 digest",
            )
        if self._source_path is None:
            return ResultEnvelope.fail(
                "NO_PROTECTED_SOURCE",
                "No source project is associated with the active working-copy session",
            )
        requested_source = Path(source_path).expanduser().resolve()
        source = self._source_path.resolve()
        if not _same_file_path(requested_source, source):
            return ResultEnvelope.fail(
                "SOURCE_PATH_MISMATCH",
                "source_path must identify the current working copy's protected source",
            )
        if source.suffix.lower() != ".opju":
            return ResultEnvelope.fail("UNSUPPORTED_PROJECT", "Only .opju source projects can be replaced")
        try:
            before = _project_file_metadata(source)
        except OSError as exc:
            return ResultEnvelope.fail(_exception_error_code(exc), str(exc))
        if before["sha256"] != expected_digest:
            return ResultEnvelope.fail(
                "SOURCE_CHANGED",
                "The source project SHA-256 no longer matches the explicitly authorized version",
                data={"source_path": str(source), "actual": before, "expected_sha256": expected_digest},
            )

        token = uuid4().hex[:12]
        candidate = source.with_name(f".{source.stem}.codex-candidate-{token}{source.suffix}")
        backup = source.with_name(f".{source.stem}.codex-backup-{token}{source.suffix}")

        def prepare_candidate() -> dict[str, Any]:
            app = self._require_app()
            fingerprint_before = _project_structure_fingerprint(app)
            run = getattr(app, "Run", None)
            if callable(run) and run() is False:
                raise ProjectSaveUnconfirmedError(
                    "Origin did not finish pending recalculation before candidate save"
                )
            saved = app.Save(str(candidate))
            if saved not in (True, 1):
                raise ProjectSaveUnconfirmedError(
                    f"Origin did not explicitly confirm the candidate save: {candidate}"
                )
            if not validate_project_artifact(candidate):
                raise ProjectSaveUnconfirmedError(
                    f"Origin did not create a plausibly sized candidate project: {candidate}"
                )
            modified = _safe_attr(app, "IsModified")
            if modified not in (None, False, 0):
                raise ProjectSaveUnconfirmedError(
                    "Origin still reports unsaved changes after candidate save"
                )
            if app.NewProject() not in (True, 1):
                raise ProjectCloseUnconfirmedError(
                    "Origin could not release the saved candidate before verification"
                )
            self._project_path = None
            try:
                loaded = app.Load(str(candidate), False)
                if loaded not in (True, 1):
                    raise ProjectSaveUnconfirmedError(
                        f"Origin could not reopen the saved candidate: {candidate}"
                    )
                fingerprint_after = _project_structure_fingerprint(app)
                if fingerprint_after != fingerprint_before:
                    raise ProjectSaveUnconfirmedError(
                        "The reopened candidate project structure differs from the active project"
                    )
            finally:
                if app.NewProject() not in (True, 1):
                    raise ProjectCloseUnconfirmedError(
                        "Origin could not release the verified candidate project"
                    )
                self._project_path = None
            return {
                "candidate": _project_file_metadata(candidate),
                "fingerprint": fingerprint_after,
            }

        prepared = self._submit(prepare_candidate)
        if not prepared.success:
            failed_artifacts = (
                [Artifact(str(candidate), "origin_project_candidate")]
                if candidate.is_file()
                else []
            )
            return replace(prepared, artifacts=failed_artifacts)

        candidate_metadata = prepared.data["candidate"]
        artifacts: list[Artifact] = []

        def rollback_failed_commit(
            reason: str,
            *,
            after: dict[str, Any] | None = None,
        ) -> ResultEnvelope:
            recovery_artifacts: list[Artifact] = []
            recovery_warnings: list[str] = []
            try:
                try:
                    shutil.copy2(source, candidate)
                except OSError as preserve_error:
                    recovery_warnings.append(
                        f"Could not preserve the failed replacement candidate: {preserve_error}"
                    )
                os.replace(backup, source)
            except OSError as rollback_error:
                if candidate.is_file():
                    recovery_artifacts.append(
                        Artifact(str(candidate), "origin_project_candidate")
                    )
                if backup.is_file():
                    recovery_artifacts.append(
                        Artifact(str(backup), "origin_project_backup")
                    )
                return ResultEnvelope.fail(
                    "SOURCE_REPLACE_ROLLBACK_FAILED",
                    f"{reason}; rollback also failed: {rollback_error}",
                    data={
                        "source_path": str(source),
                        "candidate": candidate_metadata,
                        "before": before,
                        "after": after,
                        "backup_path": str(backup),
                    },
                    warnings=recovery_warnings,
                    artifacts=recovery_artifacts,
                    duration_ms=prepared.duration_ms,
                )
            restored_signature = project_signature(source)
            if restored_signature is None or restored_signature[2] != before["sha256"]:
                if candidate.is_file():
                    recovery_artifacts.append(
                        Artifact(str(candidate), "origin_project_candidate")
                    )
                return ResultEnvelope.fail(
                    "SOURCE_REPLACE_ROLLBACK_FAILED",
                    f"{reason}; rollback returned but the authorized source hash was not restored",
                    data={
                        "source_path": str(source),
                        "candidate": candidate_metadata,
                        "before": before,
                        "after": after,
                    },
                    warnings=recovery_warnings,
                    artifacts=recovery_artifacts,
                    duration_ms=prepared.duration_ms,
                )
            if candidate.is_file():
                recovery_artifacts.append(
                    Artifact(str(candidate), "origin_project_candidate")
                )
            return ResultEnvelope.fail(
                "SOURCE_REPLACE_VALIDATION_FAILED",
                f"{reason}; the authorized source was restored from its verified backup",
                data={
                    "source_path": str(source),
                    "candidate": candidate_metadata,
                    "before": before,
                    "after": after,
                    "rollback_confirmed": True,
                },
                warnings=recovery_warnings,
                artifacts=recovery_artifacts,
                duration_ms=prepared.duration_ms,
            )

        try:
            current = _project_file_metadata(source)
            if current != before:
                raise SourceOverwriteError(
                    "The source project changed after candidate verification; replacement was cancelled"
                )
            shutil.copy2(source, backup)
            backup_metadata = _project_file_metadata(backup)
            if backup_metadata["sha256"] != before["sha256"]:
                raise OSError("The source backup SHA-256 does not match the authorized source")
            if _project_file_metadata(source) != before:
                raise SourceOverwriteError(
                    "The source project changed while its backup was being prepared"
                )
            os.replace(candidate, source)
            try:
                after = _project_file_metadata(source)
            except OSError as verification_error:
                return rollback_failed_commit(
                    f"Replacement metadata validation failed: {verification_error}"
                )
            if after["sha256"] != candidate_metadata["sha256"]:
                return rollback_failed_commit(
                    "Replacement SHA-256 did not match the verified candidate",
                    after=after,
                )
            new_identity = _file_identity(source)
            if new_identity is not None:
                self._protected_source_identities.add(new_identity)
            backup_identity = _file_identity(backup)
            if backup_identity is not None:
                self._protected_source_identities.add(backup_identity)
            artifacts.append(Artifact(str(source), "origin_project"))
            if keep_backup:
                artifacts.append(Artifact(str(backup), "origin_project_backup"))
            else:
                try:
                    backup.unlink()
                except OSError as cleanup_error:
                    artifacts.append(Artifact(str(backup), "origin_project_backup"))
                    cleanup_warning = (
                        f"Replacement succeeded but the backup could not be removed: {cleanup_error}"
                    )
                else:
                    cleanup_warning = None
            return ResultEnvelope.ok(
                {
                    "source_path": str(source),
                    "candidate_path": str(candidate),
                    "backup_path": str(backup) if backup.is_file() else None,
                    "before": before,
                    "after": after,
                    "candidate": candidate_metadata,
                    "candidate_reopen_verified": True,
                    "structural_fingerprint_match": True,
                    "atomic_replace": True,
                    "replacement_method": "os.replace_same_directory",
                    "active_project_open": False,
                },
                warnings=[
                    "The verified candidate replaced the authorized source; Origin is now on a blank project",
                    *([cleanup_warning] if not keep_backup and cleanup_warning else []),
                ],
                artifacts=artifacts,
                duration_ms=prepared.duration_ms,
            )
        except SourceOverwriteError as exc:
            if candidate.is_file():
                artifacts.append(Artifact(str(candidate), "origin_project_candidate"))
            if backup.is_file():
                artifacts.append(Artifact(str(backup), "origin_project_backup"))
            return ResultEnvelope.fail(
                "SOURCE_CHANGED",
                str(exc),
                data={"source_path": str(source), "before": before},
                artifacts=artifacts,
            )
        except OSError as exc:
            if candidate.is_file():
                artifacts.append(Artifact(str(candidate), "origin_project_candidate"))
            if backup.is_file():
                artifacts.append(Artifact(str(backup), "origin_project_backup"))
            return ResultEnvelope.fail(
                _exception_error_code(exc),
                str(exc),
                data={"source_path": str(source), "before": before},
                artifacts=artifacts,
            )

    @_serialized_operation
    def list_objects(self) -> ResultEnvelope:
        def describe() -> dict[str, Any]:
            app = self._require_app()
            worksheet_catalog = _worksheet_column_catalog(app)
            pages: list[dict[str, Any]] = []
            worksheets: list[dict[str, Any]] = []
            graph_layers: list[dict[str, Any]] = []
            matrix_sheets: list[dict[str, Any]] = []
            data_sources: list[dict[str, Any]] = []
            seen: set[str] = set()
            typed = (
                ("WorksheetPages", "worksheet_page"),
                ("GraphPages", "graph_page"),
                ("MatrixPages", "matrix_page"),
            )
            for collection_name, kind in typed:
                for page in _required_collection_items(_required_attr(app, collection_name)):
                    name = str(_safe_attr(page, "Name", ""))
                    long_name = str(_safe_attr(page, "LongName", name))
                    key = name or long_name
                    if key in seen:
                        continue
                    seen.add(key)
                    layers: list[dict[str, Any]] = []
                    for layer_position, raw_layer in enumerate(
                        _required_collection_items(_required_attr(page, "Layers"))
                    ):
                        layer_name = str(_safe_attr(raw_layer, "Name", ""))
                        layer = raw_layer
                        try:
                            if kind == "graph_page":
                                layer = app.FindGraphLayer(f"[{name}]{layer_position + 1}") or raw_layer
                            elif kind == "worksheet_page" and layer_name:
                                layer = app.FindWorksheet(f"[{name}]{layer_name}") or raw_layer
                            elif kind == "matrix_page" and layer_name:
                                layer = app.FindMatrixSheet(f"[{name}]{layer_name}") or raw_layer
                        except Exception:
                            layer = raw_layer
                        if kind == "graph_page":
                            layer_reference = f"[{name}]{layer_position + 1}"
                        else:
                            layer_reference = f"[{name}]{layer_name}"
                        if kind == "worksheet_page":
                            column_collection = _required_attr(layer, "Columns") or _required_attr(
                                layer, "DataObjectBases"
                            )
                            column_items = _required_collection_items(column_collection)
                        else:
                            column_collection = _safe_attr(layer, "Columns") or _safe_attr(
                                layer, "DataObjectBases"
                            )
                            column_items = _collection_items(column_collection)
                        columns = [
                            {
                                "index": _safe_attr(column, "Index"),
                                "name": str(_safe_attr(column, "Name", "")),
                                "long_name": str(_safe_attr(column, "LongName", "")),
                                "units": str(_safe_attr(column, "Units", "")),
                                "comments": str(_safe_attr(column, "Comments", "")),
                                "plot_designation": _safe_attr(column, "PlotDesignation"),
                            }
                            for column in column_items
                        ]
                        layer_descriptor = {
                            "id": layer_reference,
                            "ref": layer_reference,
                            "page_name": name,
                            "name": layer_name,
                            "long_name": str(_safe_attr(layer, "LongName", "")),
                            "index": _safe_attr(layer, "Index"),
                            "rows": _safe_attr(layer, "Rows"),
                            "cols": _safe_attr(layer, "Cols"),
                            "columns": columns,
                        }
                        layers.append(layer_descriptor)
                        if kind == "worksheet_page":
                            worksheets.append(layer_descriptor)
                        elif kind == "graph_page":
                            graph_layers.append(layer_descriptor)
                        elif kind == "matrix_page":
                            matrix_sheets.append(layer_descriptor)
                        for source_collection in ("DataPlots", "Columns", "DataObjectBases"):
                            if source_collection == "DataPlots" and kind == "graph_page":
                                source_items = _required_collection_items(
                                    _required_attr(layer, source_collection)
                                )
                            else:
                                source_items = _collection_items(
                                    _safe_attr(layer, source_collection)
                                )
                            for source_position, source in enumerate(source_items, start=1):
                                source_name = str(_safe_attr(source, "Name", ""))
                                source_descriptor = {
                                    "id": f"{layer_reference}:{source_collection}:{source_name}",
                                    "page_name": name,
                                    "layer_name": layer_name,
                                    "collection": source_collection,
                                    "name": source_name,
                                    "type": str(_safe_attr(source, "TypeName", "")),
                                    "range": str(_safe_attr(source, "Range", "")),
                                }
                                if source_collection == "DataPlots":
                                    source_descriptor = _describe_data_plot(
                                        app,
                                        layer,
                                        source,
                                        source_descriptor,
                                        graph_name=name,
                                        layer_position=layer_position + 1,
                                        plot_index=source_position,
                                        catalog=worksheet_catalog,
                                    )
                                data_sources.append(source_descriptor)
                    pages.append(
                        {
                            "id": f"{kind}:{name or long_name}",
                            "ref": name,
                            "kind": kind,
                            "name": name,
                            "long_name": long_name,
                            "type": str(_safe_attr(page, "Type", _safe_attr(page, "TypeName", ""))),
                            "pe_path": str(_safe_attr(page, "PEPath", "")),
                            "layers": layers,
                        }
                    )
            return {
                "pages": pages,
                "workbooks": [page for page in pages if page["kind"] == "worksheet_page"],
                "worksheets": worksheets,
                "graphs": [page for page in pages if page["kind"] == "graph_page"],
                "graph_layers": graph_layers,
                "matrices": [page for page in pages if page["kind"] == "matrix_page"],
                "matrix_sheets": matrix_sheets,
                "data_sources": data_sources,
            }

        return self._submit(describe, retryable=True)

    @_serialized_operation
    def transform_worksheet(
        self,
        *,
        source_ref: str,
        destination_ref: str,
        action: str,
        options: Mapping[str, Any] | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        if source_ref == destination_ref:
            return ResultEnvelope.fail(
                "IN_PLACE_TRANSFORM_BLOCKED",
                "Worksheet transforms require a distinct destination_ref",
            )

        def execute() -> ResultEnvelope:
            app = self._require_app()
            source = _resolve_worksheet(app, source_ref)
            destination = _resolve_worksheet(app, destination_ref)
            source_rows = table_from_com_value(
                source.GetData(0, 0, -1, -1, ORIGIN_ARRAY2D_VARIANT)
            )
            column_items = _collection_items(_safe_attr(source, "Columns"))
            columns = [
                str(_safe_attr(item, "LongName", "") or _safe_attr(item, "Name", ""))
                for item in column_items
            ]
            try:
                transformed = transform_table(
                    source_rows,
                    columns,
                    action=action,
                    options=dict(options or {}),
                )
            except TransformValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))
            destination.Cols = len(transformed.columns)
            set_result = destination.SetData(transformed.rows, 0, 0)
            if set_result is False or set_result == 0:
                return ResultEnvelope.fail(
                    "WORKSHEET_TRANSFORM_WRITE_REJECTED",
                    "Origin rejected the transformed worksheet block",
                )
            destination_columns = _collection_items(_safe_attr(destination, "Columns"))
            for index, label in enumerate(transformed.columns):
                if index < len(destination_columns):
                    try:
                        destination_columns[index].LongName = label
                    except Exception:
                        pass
            readback = _readback_matrix(
                destination,
                row=0,
                column=0,
                rows=len(transformed.rows),
                columns=len(transformed.columns),
            )
            mismatches = _matrix_mismatches(transformed.rows, readback)
            data = {
                "action": transformed.action,
                "source_ref": source_ref,
                "destination_ref": destination_ref,
                "input_rows": transformed.input_rows,
                "output_rows": transformed.output_rows,
                "columns": transformed.columns,
                "readback_verified": not mismatches,
                "mismatches": mismatches,
            }
            if mismatches:
                return ResultEnvelope.fail(
                    "WORKSHEET_TRANSFORM_UNCONFIRMED",
                    "Transformed worksheet readback did not match",
                    data=data,
                )
            return ResultEnvelope.ok(data)

        return self._submit(execute, stage=f"transform_worksheet_{action}")

    @_serialized_operation
    def manage_connector(
        self,
        *,
        action: str,
        worksheet_ref: str,
        source: str | None = None,
        connector_type: str | None = None,
        keep_connector: bool = True,
        keep_data: bool | None = None,
    ) -> ResultEnvelope:
        if action != "info":
            guard = self._guard_mutation()
            if guard:
                return guard
        try:
            plan = build_connector_plan(
                action=action,
                worksheet_ref=worksheet_ref,
                source=source,
                connector_type=connector_type,
                keep_connector=keep_connector,
                keep_data=keep_data,
            )
        except ObjectPlanError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))

        def execute() -> ResultEnvelope:
            app = self._require_app()
            sheet = _resolve_worksheet(app, plan.worksheet_ref)
            connector = _safe_attr(sheet, "Connector")
            if connector is None:
                required_methods = (
                    "DoMethod",
                    "Execute",
                    "GetNumProp",
                    "GetStrProp",
                    "SetStrProp",
                )
                if not all(callable(getattr(sheet, name, None)) for name in required_methods):
                    return ResultEnvelope.fail(
                        "CONNECTOR_INTERFACE_UNAVAILABLE",
                        "This Origin worksheet does not expose Data Connector COM methods",
                    )
                connected = bool(sheet.GetNumProp("HasDC"))
                if plan.action == "create":
                    sheet.DoMethod("DC.Allow", "2")
                    add_result = sheet.Execute(
                        f"wbook.dc.add({plan.connector_type})"
                    )
                    if add_result is False or add_result == 0:
                        return ResultEnvelope.fail(
                            "CONNECTOR_ACTION_REJECTED",
                            "Origin rejected Data Connector creation",
                        )
                    sheet.SetStrProp("DC.Source", str(plan.source))
                    previous_sparklines = _safe_call(
                        app, "LTVar", "@IMPS", default=None
                    )
                    _safe_call(app, "Execute", "@IMPS=0;", default=None)
                    try:
                        sheet.DoMethod("DC.Import", "")
                    finally:
                        if previous_sparklines is not None:
                            _safe_call(
                                app,
                                "Execute",
                                f"@IMPS={int(previous_sparklines)};",
                                default=None,
                            )
                    connected = bool(sheet.GetNumProp("HasDC"))
                elif plan.action == "refresh":
                    if not connected:
                        return ResultEnvelope.fail(
                            "CONNECTOR_NOT_FOUND",
                            f"Worksheet has no Data Connector: {plan.worksheet_ref}",
                        )
                    previous_sparklines = _safe_call(
                        app, "LTVar", "@IMPS", default=None
                    )
                    _safe_call(app, "Execute", "@IMPS=0;", default=None)
                    try:
                        sheet.DoMethod("DC.Import", "")
                    finally:
                        if previous_sparklines is not None:
                            _safe_call(
                                app,
                                "Execute",
                                f"@IMPS={int(previous_sparklines)};",
                                default=None,
                            )
                elif plan.action == "disconnect":
                    if not connected:
                        return ResultEnvelope.fail(
                            "CONNECTOR_NOT_FOUND",
                            f"Worksheet has no Data Connector: {plan.worksheet_ref}",
                        )
                    remove_result = sheet.Execute("wbook.dc.remove();")
                    if remove_result is False or remove_result == 0:
                        return ResultEnvelope.fail(
                            "CONNECTOR_ACTION_REJECTED",
                            "Origin rejected Data Connector removal",
                        )
                    connected = bool(sheet.GetNumProp("HasDC"))
                actual_source = str(sheet.GetStrProp("DC.Source") or "")
                if (
                    plan.action == "create"
                    and Path(actual_source).resolve() != plan.source
                ):
                    return ResultEnvelope.fail(
                        "CONNECTOR_CREATE_UNCONFIRMED",
                        "Data Connector source readback did not match",
                        data={"source": actual_source, "connected": connected},
                    )
                if plan.action == "create" and not connected:
                    return ResultEnvelope.fail(
                        "CONNECTOR_CREATE_UNCONFIRMED",
                        "Worksheet did not report an active Data Connector",
                    )
                if plan.action == "disconnect" and connected:
                    return ResultEnvelope.fail(
                        "CONNECTOR_DISCONNECT_UNCONFIRMED",
                        "Data Connector remained active after removal",
                    )
                parent = _safe_attr(sheet, "Parent")
                page_name = str(_safe_attr(parent, "Name", "") or "")
                sheet_name = str(_safe_attr(sheet, "Name", "") or "")
                actual_worksheet_ref = (
                    f"[{page_name}]{sheet_name}"
                    if page_name and sheet_name
                    else plan.worksheet_ref
                )
                connector_type_value = str(
                    _safe_call(parent, "GetStrProp", "DC.Type", default="") or ""
                )
                return ResultEnvelope.ok(
                    {
                        "action": plan.action,
                        "requested_worksheet_ref": plan.worksheet_ref,
                        "worksheet_ref": actual_worksheet_ref,
                        "source": actual_source or None,
                        "connector_type": connector_type_value or plan.connector_type,
                        "connected": connected,
                        "refresh_count": int(_safe_attr(sheet, "refreshes", 0) or 0),
                        "keep_data": plan.keep_data,
                        "interface": "labtalk_data_connector",
                    }
                )
            if plan.action == "create":
                result = connector.Connect(
                    str(plan.source), plan.connector_type, plan.keep_connector
                )
            elif plan.action == "refresh":
                result = connector.Refresh()
            elif plan.action == "disconnect":
                result = connector.Disconnect(bool(plan.keep_data))
            else:
                result = True
            if result is False or result == 0:
                return ResultEnvelope.fail(
                    "CONNECTOR_ACTION_REJECTED",
                    f"Origin rejected connector {plan.action}",
                )
            connected = bool(
                _safe_attr(connector, "Connected", _safe_attr(connector, "connected", False))
            )
            actual_source = str(
                _safe_attr(connector, "Source", _safe_attr(connector, "source", ""))
            )
            if plan.action == "create" and Path(actual_source).resolve() != plan.source:
                return ResultEnvelope.fail(
                    "CONNECTOR_CREATE_UNCONFIRMED",
                    "Connector source readback did not match",
                    data={"source": actual_source, "connected": connected},
                )
            if plan.action == "disconnect" and connected:
                return ResultEnvelope.fail(
                    "CONNECTOR_DISCONNECT_UNCONFIRMED",
                    "Connector remained connected after disconnect",
                )
            return ResultEnvelope.ok(
                {
                    "action": plan.action,
                    "worksheet_ref": plan.worksheet_ref,
                    "source": actual_source or None,
                    "connector_type": str(
                        _safe_attr(connector, "Type", _safe_attr(connector, "type", ""))
                    ) or None,
                    "connected": connected,
                    "refresh_count": int(_safe_attr(connector, "refreshes", 0) or 0),
                    "keep_data": plan.keep_data,
                }
            )

        return self._submit(
            execute,
            retryable=plan.action == "info",
            stage=f"connector_{plan.action}",
        )

    @_serialized_operation
    def manage_matrix(
        self,
        *,
        action: str,
        matrix_ref: str,
        values: list[list[Any]] | None = None,
        row: int = 0,
        column: int = 0,
        operation: str | None = None,
    ) -> ResultEnvelope:
        if action != "read":
            guard = self._guard_mutation()
            if guard:
                return guard
        try:
            plan = build_matrix_plan(
                action=action,
                matrix_ref=matrix_ref,
                values=values,
                row=row,
                column=column,
                operation=operation,
            )
        except ObjectPlanError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))

        def execute() -> ResultEnvelope:
            app = self._require_app()
            if plan.action == "create":
                pages = _safe_attr(app, "MatrixPages")
                if pages is None or not callable(getattr(pages, "Add", None)):
                    return ResultEnvelope.fail(
                        "MATRIX_CREATE_UNAVAILABLE",
                        "Origin does not expose MatrixPages.Add through COM",
                    )
                before_count = int(_safe_attr(pages, "Count", 0) or 0)
                page = pages.Add()
                if page is None:
                    return ResultEnvelope.fail(
                        "MATRIX_CREATE_UNCONFIRMED", "MatrixPages.Add returned no page"
                    )
                try:
                    page.LongName = plan.matrix_ref
                except Exception:
                    pass
                layers = _safe_attr(page, "Layers")
                layer_items = _collection_items(layers)
                after_count = int(_safe_attr(pages, "Count", 0) or 0)
                if after_count != before_count + 1 or not layer_items:
                    return ResultEnvelope.fail(
                        "MATRIX_CREATE_UNCONFIRMED",
                        "Matrix page or first matrix sheet was not created",
                    )
                layer = layer_items[0]
                actual_ref = (
                    f'[{_safe_attr(page, "Name", "")}]'
                    f'{_safe_attr(layer, "Name", "")}'
                )
                matrix_sheet = _safe_call(
                    app, "FindMatrixSheet", actual_ref, default=None
                )
                matrix_objects = _safe_attr(matrix_sheet, "MatrixObjects")
                if (
                    matrix_objects is not None
                    and int(_safe_attr(matrix_objects, "Count", 0) or 0) == 0
                    and callable(getattr(matrix_objects, "Add", None))
                ):
                    matrix_objects.Add()
                data_object = _matrix_data_object(matrix_sheet)
                if matrix_sheet is None or not all(
                    callable(getattr(data_object, method, None))
                    for method in ("SetData", "GetData")
                ):
                    return ResultEnvelope.fail(
                        "MATRIX_CREATE_UNCONFIRMED",
                        "Matrix page was created without a writable matrix object",
                    )
                return ResultEnvelope.ok(
                    {
                        "action": "create",
                        "requested_ref": plan.matrix_ref,
                        "matrix_ref": actual_ref,
                        "page_name": str(_safe_attr(page, "Name", "")),
                        "sheet_name": str(_safe_attr(layer, "Name", "")),
                        "verified": True,
                    }
                )
            matrix = _safe_call(app, "FindMatrixSheet", plan.matrix_ref, default=None)
            if matrix is None:
                return ResultEnvelope.fail(
                    "MATRIX_NOT_FOUND", f"Matrix sheet not found: {plan.matrix_ref}"
                )
            matrix_object = _matrix_data_object(matrix)
            column_major_bridge = matrix_object is not matrix
            if plan.action == "write":
                block = [list(item) for item in plan.values or ()]
                result = matrix_object.SetData(block, plan.row, plan.column)
                if result is False or result == 0:
                    return ResultEnvelope.fail("MATRIX_WRITE_REJECTED", "Origin rejected Matrix.SetData")
                raw = _matrix_object_block(
                    matrix_object,
                    plan.row,
                    plan.column,
                    len(block),
                    len(block[0]) if block else 0,
                )
                readback = _normalize_matrix_object_table(
                    raw, column_major_bridge=column_major_bridge
                )
                mismatches = _matrix_mismatches(block, readback)
                if mismatches:
                    return ResultEnvelope.fail(
                        "MATRIX_WRITE_UNCONFIRMED",
                        "Matrix readback did not match",
                        data={"mismatches": mismatches},
                    )
                return ResultEnvelope.ok(
                    {
                        "action": "write",
                        "matrix_ref": plan.matrix_ref,
                        "shape": list(plan.shape or (0, 0)),
                        "expected_range": list(plan.expected_range or ()),
                        "readback_verified": True,
                    }
                )
            if plan.action == "read":
                data = _normalize_matrix_object_table(
                    matrix_object.GetData(),
                    column_major_bridge=column_major_bridge,
                )
                return ResultEnvelope.ok(
                    {
                        "action": "read",
                        "matrix_ref": plan.matrix_ref,
                        "values": data,
                        "shape": [len(data), len(data[0]) if data else 0],
                    }
                )
            if plan.action == "transform":
                result = matrix.Execute(plan.command)
                if result is False or result == 0:
                    return ResultEnvelope.fail("MATRIX_TRANSFORM_REJECTED", "Origin rejected matrix transform")
                return ResultEnvelope.ok(
                    {
                        "action": "transform",
                        "matrix_ref": plan.matrix_ref,
                        "operation": plan.operation,
                        "supported_unverified": True,
                    }
                )
            return ResultEnvelope.fail("MATRIX_ACTION_UNSUPPORTED", plan.action)

        return self._submit(
            execute, retryable=plan.action == "read", stage=f"matrix_{plan.action}"
        )

    @_serialized_operation
    def manage_image(
        self,
        *,
        action: str,
        image_ref: str,
        path: str | None = None,
        overwrite: bool = False,
    ) -> ResultEnvelope:
        if action != "info":
            guard = self._guard_mutation()
            if guard:
                return guard
        try:
            plan = build_image_plan(
                action=action,
                image_ref=image_ref,
                path=path,
                overwrite=overwrite,
            )
        except ObjectPlanError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))

        def execute() -> ResultEnvelope:
            app = self._require_app()
            page = _safe_call(app, "FindImagePage", plan.image_ref, default=None)
            pages = _safe_attr(app, "ImagePages")
            if page is None:
                page = _find_collection_item(pages, plan.image_ref)
            if plan.action == "create":
                if page is not None:
                    return ResultEnvelope.fail(
                        "IMAGE_PAGE_ALREADY_EXISTS",
                        f"Image Page already exists: {plan.image_ref}",
                    )
                if pages is None or not callable(getattr(pages, "Add", None)):
                    return ResultEnvelope.fail(
                        "IMAGE_CREATE_UNAVAILABLE",
                        "Origin does not expose ImagePages.Add through COM",
                    )
                before_count = int(_safe_attr(pages, "Count", 0) or 0)
                page = pages.Add()
                if page is None:
                    return ResultEnvelope.fail(
                        "IMAGE_CREATE_UNCONFIRMED", "ImagePages.Add returned no page"
                    )
                try:
                    page.LongName = plan.image_ref
                except Exception:
                    pass
                if int(_safe_attr(pages, "Count", 0) or 0) != before_count + 1:
                    return ResultEnvelope.fail(
                        "IMAGE_CREATE_UNCONFIRMED", "ImagePages count did not increase"
                    )
                return ResultEnvelope.ok(
                    {
                        "action": "create",
                        "requested_ref": plan.image_ref,
                        "image_ref": str(_safe_attr(page, "Name", "")),
                        "long_name": str(_safe_attr(page, "LongName", "")),
                        "verified": True,
                    }
                )
            if page is None:
                return ResultEnvelope.fail("IMAGE_PAGE_NOT_FOUND", f"Image Page not found: {plan.image_ref}")
            artifacts: list[Artifact] = []
            if plan.action == "import":
                importer = getattr(page, "Import", None)
                if callable(importer):
                    result = importer(str(plan.path))
                else:
                    result = page.Execute(
                        f"img.Load({_labtalk_quote(str(plan.path))})"
                    )
            elif plan.action == "export":
                exporter = getattr(page, "Export", None)
                if callable(exporter):
                    result = exporter(str(plan.path))
                else:
                    result = page.Preview(str(plan.path))
            elif plan.action == "delete":
                delete = getattr(page, "Delete", None) or getattr(page, "Destroy", None)
                if not callable(delete):
                    return ResultEnvelope.fail("IMAGE_DELETE_UNAVAILABLE", "Image Page does not expose Delete")
                result = delete()
            else:
                result = True
            if result is False or result == 0:
                return ResultEnvelope.fail("IMAGE_ACTION_REJECTED", f"Origin rejected image {plan.action}")
            if plan.action == "export":
                if not plan.path or not plan.path.is_file() or plan.path.stat().st_size == 0:
                    return ResultEnvelope.fail("IMAGE_EXPORT_UNCONFIRMED", "Image export artifact is missing or empty")
                artifacts.append(Artifact(str(plan.path), "origin_image"))
            width = _page_numeric_property(page, "Width")
            height = _page_numeric_property(page, "Height")
            source_value = str(_safe_attr(page, "Source", "")) or None
            if plan.action == "import":
                source_matches = bool(
                    source_value and Path(source_value).resolve() == plan.path
                )
                if not source_matches and not (width and height):
                    return ResultEnvelope.fail(
                        "IMAGE_IMPORT_UNCONFIRMED",
                        "Image source and dimensions could not be confirmed",
                    )
                if not source_value:
                    source_value = str(plan.path)
            return ResultEnvelope.ok(
                {
                    "action": plan.action,
                    "image_ref": plan.image_ref,
                    "source": source_value,
                    "width": width,
                    "height": height,
                    "path": str(plan.path) if plan.path else None,
                    "artifact_verified": plan.action != "export" or bool(artifacts),
                },
                artifacts=artifacts,
            )

        return self._submit(
            execute, retryable=plan.action == "info", stage=f"image_{plan.action}"
        )

    @_serialized_operation
    def manage_project_folder(
        self,
        *,
        action: str,
        path: str,
        destination: str | None = None,
        confirm_recursive: bool = False,
    ) -> ResultEnvelope:
        if action != "list":
            guard = self._guard_mutation()
            if guard:
                return guard
        try:
            plan = build_folder_plan(
                action=action,
                path=path,
                destination=destination,
                confirm_recursive=confirm_recursive,
            )
        except ProjectObjectError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))

        def execute() -> ResultEnvelope:
            app = self._require_app()
            folders = _safe_attr(app, "ProjectFolders")
            if folders is None:
                root = _safe_attr(app, "RootFolder")
                if root is None:
                    return ResultEnvelope.fail(
                        "PROJECT_FOLDER_INTERFACE_UNAVAILABLE",
                        "This Origin version exposes neither ProjectFolders nor RootFolder",
                    )
                if plan.action == "list":
                    folder = _resolve_root_folder(root, plan.path)
                    if folder is None:
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_NOT_FOUND",
                            f"Project Folder not found: {plan.path}",
                        )
                    prefix = plan.path.rstrip("/")
                    children = [
                        f'{prefix}/{_safe_attr(child, "Name", "")}'
                        for child in _collection_items(_safe_attr(folder, "Folders"))
                    ]
                    return ResultEnvelope.ok(
                        {"action": "list", "path": plan.path, "children": children}
                    )
                if plan.action == "create":
                    folder = _ensure_root_folder(root, plan.path)
                    verified = folder is not None
                elif plan.action == "rename":
                    folder = _resolve_root_folder(root, plan.path)
                    source_parent = "/" + "/".join(_project_path_parts(plan.path)[:-1])
                    destination_parent = "/" + "/".join(
                        _project_path_parts(plan.destination or "")[:-1]
                    )
                    if folder is None:
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_NOT_FOUND",
                            f"Project Folder not found: {plan.path}",
                        )
                    if source_parent.rstrip("/") != destination_parent.rstrip("/"):
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_MOVE_UNAVAILABLE",
                            "RootFolder COM supports rename only within the current parent",
                        )
                    new_name = _project_path_parts(plan.destination or "")[-1]
                    folder.Name = new_name
                    try:
                        folder.LongName = new_name
                    except Exception:
                        pass
                    verified = _resolve_root_folder(root, plan.destination or "") is not None
                elif plan.action == "move":
                    return ResultEnvelope.fail(
                        "PROJECT_FOLDER_MOVE_UNAVAILABLE",
                        "RootFolder COM does not expose a verified folder move method",
                    )
                else:
                    folder = _resolve_root_folder(root, plan.path)
                    if folder is None:
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_NOT_FOUND",
                            f"Project Folder not found: {plan.path}",
                        )
                    children = _collection_items(_safe_attr(folder, "Folders"))
                    pages = _collection_items(_safe_attr(folder, "PageBases"))
                    if (children or pages) and not plan.confirm_recursive:
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_NOT_EMPTY",
                            "Non-empty folder deletion requires confirm_recursive=true",
                        )
                    destroy = getattr(folder, "Destroy", None)
                    if not callable(destroy):
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_DELETE_UNAVAILABLE",
                            "Folder does not expose Destroy through COM",
                        )
                    result = destroy()
                    if result is False or result == 0:
                        return ResultEnvelope.fail(
                            "PROJECT_FOLDER_ACTION_REJECTED",
                            "Origin rejected folder deletion",
                        )
                    verified = _resolve_root_folder(
                        _safe_attr(app, "RootFolder"), plan.path
                    ) is None
                data = {
                    "action": plan.action,
                    "path": plan.path,
                    "destination": plan.destination,
                    "verified": verified,
                    "interface": "root_folder_com",
                }
                if not verified:
                    return ResultEnvelope.fail(
                        "PROJECT_FOLDER_ACTION_UNCONFIRMED",
                        "RootFolder state did not confirm the requested action",
                        data=data,
                    )
                return ResultEnvelope.ok(data)
            def exists(value: str) -> bool:
                return bool(_safe_call(folders, "Exists", value, default=False))
            if plan.action == "list":
                children = _safe_call(folders, "List", plan.path, default=None)
                if children is None:
                    return ResultEnvelope.fail("PROJECT_FOLDER_LIST_UNAVAILABLE", "ProjectFolders.List is unavailable")
                return ResultEnvelope.ok(
                    {"action": "list", "path": plan.path, "children": list(children)}
                )
            if plan.action == "create":
                result = folders.Create(plan.path)
                verified = exists(plan.path)
            elif plan.action in {"move", "rename"}:
                result = folders.Move(plan.path, plan.destination)
                verified = bool(plan.destination and exists(plan.destination) and not exists(plan.path))
            else:
                children = list(_safe_call(folders, "List", plan.path, default=[]))
                if children and not plan.confirm_recursive:
                    return ResultEnvelope.fail(
                        "PROJECT_FOLDER_NOT_EMPTY",
                        "Non-empty folder deletion requires confirm_recursive=true",
                        data={"children": children},
                    )
                result = folders.Delete(plan.path, plan.confirm_recursive)
                verified = not exists(plan.path)
            if result is False or result == 0:
                return ResultEnvelope.fail(
                    "PROJECT_FOLDER_ACTION_REJECTED",
                    f"Origin rejected folder {plan.action}",
                )
            data = {
                "action": plan.action,
                "path": plan.path,
                "destination": plan.destination,
                "verified": verified,
            }
            if not verified:
                return ResultEnvelope.fail(
                    "PROJECT_FOLDER_ACTION_UNCONFIRMED",
                    "Project Folder state did not confirm the requested action",
                    data=data,
                )
            return ResultEnvelope.ok(data)

        return self._submit(
            execute,
            retryable=plan.action == "list",
            stage=f"project_folder_{plan.action}",
        )

    @_serialized_operation
    def manage_note(
        self,
        *,
        action: str,
        note_ref: str,
        text: str | None = None,
        format: str = "text",
        path: str | None = None,
        overwrite: bool = False,
    ) -> ResultEnvelope:
        if action not in {"info", "export"}:
            guard = self._guard_mutation()
            if guard:
                return guard
        try:
            plan = build_note_plan(
                action=action,
                note_ref=note_ref,
                text=text,
                format=format,
                path=path,
                overwrite=overwrite,
            )
        except (ProjectObjectError, ObjectPlanError) as exc:
            return ResultEnvelope.fail(getattr(exc, "code", "PROJECT_OBJECT_INVALID"), str(exc))

        def execute() -> ResultEnvelope:
            app = self._require_app()
            note = _safe_call(app, "FindNotePage", plan.note_ref, default=None)
            notes = _safe_attr(app, "Notes")
            if note is None:
                note = _find_collection_item(notes, plan.note_ref)
            if plan.action == "create":
                if note is not None:
                    return ResultEnvelope.fail("NOTE_ALREADY_EXISTS", f"Note already exists: {plan.note_ref}")
                note = _safe_call(app, "CreateNotePage", plan.note_ref, default=None)
                if note is None and notes is not None and callable(getattr(notes, "Add", None)):
                    note = notes.Add()
                    if note is not None:
                        try:
                            note.LongName = plan.note_ref
                        except Exception:
                            pass
            if note is None:
                return ResultEnvelope.fail("NOTE_NOT_FOUND", f"Note not found: {plan.note_ref}")
            artifacts: list[Artifact] = []
            if plan.action in {"create", "write"}:
                note.Text = plan.text
                try:
                    note.Format = plan.format
                except Exception:
                    pass
                if str(_safe_attr(note, "Text", "")) != plan.text:
                    return ResultEnvelope.fail("NOTE_WRITE_UNCONFIRMED", "Note text readback did not match")
            elif plan.action == "export":
                export = getattr(note, "Export", None)
                if not callable(export):
                    return ResultEnvelope.fail("NOTE_EXPORT_UNAVAILABLE", "Note does not expose Export")
                result = export(str(plan.path), plan.format)
                if result is False or result == 0 or not plan.path or not plan.path.is_file() or plan.path.stat().st_size == 0:
                    return ResultEnvelope.fail("NOTE_EXPORT_UNCONFIRMED", "Note export artifact is missing or empty")
                artifacts.append(Artifact(str(plan.path), "origin_note"))
            elif plan.action == "delete":
                delete = getattr(note, "Delete", None)
                if not callable(delete):
                    return ResultEnvelope.fail("NOTE_DELETE_UNAVAILABLE", "Note does not expose Delete")
                result = delete()
                if result is False or result == 0:
                    return ResultEnvelope.fail("NOTE_DELETE_REJECTED", "Origin rejected Note deletion")
            return ResultEnvelope.ok(
                {
                    "action": plan.action,
                    "requested_ref": plan.note_ref,
                    "note_ref": str(_safe_attr(note, "Name", plan.note_ref)),
                    "text": str(_safe_attr(note, "Text", "")) if plan.action != "delete" else None,
                    "format": str(_safe_attr(note, "Format", plan.format)),
                    "path": str(plan.path) if plan.path else None,
                },
                artifacts=artifacts,
            )

        return self._submit(
            execute,
            retryable=plan.action == "info",
            stage=f"note_{plan.action}",
        )

    @_serialized_operation
    def write_worksheet(
        self,
        name: str,
        values: list[list[Any]],
        *,
        row: int = 0,
        column: int = 0,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        rectangular, width = _rectangular_values(values)
        if row < 0 or column < 0:
            return ResultEnvelope.fail(
                "INVALID_WORKSHEET_RANGE",
                "row and column offsets must be non-negative",
            )

        def write() -> ResultEnvelope:
            app = self._require_app()
            sheet = _resolve_worksheet(app, name)
            initial_columns = int(_safe_attr(sheet, "Cols", 0) or 0)
            if width and initial_columns < column + width:
                sheet.Cols = column + width
            formats: list[int | None] = []
            for offset in range(width):
                target_column = _worksheet_column(sheet, column + offset)
                column_values = [item[offset] for item in rectangular]
                formats.append(
                    _set_column_data_format(
                        target_column,
                        column_values,
                        force=column + offset >= initial_columns,
                    )
                )
            set_result = sheet.SetData(rectangular, row, column)
            base_data = {
                "worksheet": name,
                "worksheet_ref": name,
                "rows": len(rectangular),
                "columns": width,
                "range": [
                    row,
                    column,
                    row + len(rectangular) - 1 if rectangular else row - 1,
                    column + width - 1 if width else column - 1,
                ],
                "set_data_return": set_result,
                "column_data_formats": formats,
                "stage": "write",
            }
            if set_result is False or (
                isinstance(set_result, int)
                and not isinstance(set_result, bool)
                and set_result == 0
            ):
                return ResultEnvelope.fail(
                    "WORKSHEET_WRITE_REJECTED",
                    "Origin explicitly rejected Worksheet.SetData",
                    data=base_data,
                )
            readback = _readback_matrix(
                sheet,
                row=row,
                column=column,
                rows=len(rectangular),
                columns=width,
            )
            mismatches = _matrix_mismatches(rectangular, readback)
            verified_data = {
                **base_data,
                "stage": "verify",
                "readback_verified": not mismatches,
                "non_empty_count": sum(
                    profile["non_empty_count"]
                    for profile in (
                        _column_profile([item[offset] for item in readback])
                        for offset in range(width)
                    )
                ),
                "mismatches": mismatches,
            }
            if mismatches:
                return ResultEnvelope.fail(
                    "WORKSHEET_WRITE_UNCONFIRMED",
                    "Origin readback did not match the requested worksheet write",
                    data=verified_data,
                )
            return ResultEnvelope.ok(verified_data)

        return self._submit(write)

    @_serialized_operation
    def read_worksheet(
        self,
        name: str,
        *,
        r1: int = 0,
        c1: int = 0,
        r2: int = -1,
        c2: int = -1,
        data_format: str = "auto",
    ) -> ResultEnvelope:
        normalized_format = str(data_format).lower()
        if normalized_format not in {*WORKSHEET_DATA_FORMATS, "categorical_label"}:
            return ResultEnvelope.fail(
                "INVALID_DATA_FORMAT",
                "data_format must be auto, numeric, string, variant, or categorical_label",
            )

        def read() -> dict[str, Any]:
            app = self._require_app()
            sheet = _resolve_worksheet(app, name)
            if normalized_format == "categorical_label":
                end_column = int(_safe_attr(sheet, "Cols", 0)) - 1 if c2 == -1 else c2
                end_row = int(_safe_attr(sheet, "Rows", 0)) - 1 if r2 == -1 else r2
                columns = _safe_attr(sheet, "Columns")
                column_values = [
                    _column_string_values(columns.Item(index), r1, end_row)
                    for index in range(c1, end_column + 1)
                ]
                row_count = max((len(values) for values in column_values), default=0)
                table = [
                    [values[row] if row < len(values) else None for values in column_values]
                    for row in range(row_count)
                ]
            else:
                values = sheet.GetData(
                    r1,
                    c1,
                    r2,
                    c2,
                    WORKSHEET_DATA_FORMATS[normalized_format],
                )
                table = table_from_com_value(values)
            return {
                "worksheet": name,
                "range": [r1, c1, r2, c2],
                "data_format": normalized_format,
                "values": table,
            }

        return self._submit(read, retryable=True)

    @_serialized_operation
    def import_data(
        self,
        *,
        file_path: str,
        worksheet_name: str | None = None,
        sheet_name: str | None = None,
        has_header: bool | None = None,
        target_mode: str = "new_workbook",
        source_mode: str = "snapshot",
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        if target_mode not in {"new_workbook", "existing_worksheet"}:
            return ResultEnvelope.fail(
                "INVALID_IMPORT_TARGET_MODE",
                "target_mode must be new_workbook or existing_worksheet",
            )
        if source_mode not in {"linked", "snapshot"}:
            return ResultEnvelope.fail(
                "INVALID_IMPORT_SOURCE_MODE",
                "source_mode must be linked or snapshot",
            )
        if source_mode == "linked":
            return ResultEnvelope.fail(
                "LINKED_IMPORT_UNAVAILABLE",
                "Linked import is not available in this controller build",
            )
        source = Path(file_path).expanduser().resolve()
        if not source.is_file():
            return ResultEnvelope.fail("INPUT_NOT_FOUND", f"Data file does not exist: {source}")
        try:
            if source.suffix.lower() in {".xlsx", ".xlsm"}:
                from openpyxl import load_workbook

                workbook = load_workbook(source, read_only=True, data_only=True)
                try:
                    selected_sheet = sheet_name or workbook.sheetnames[0]
                    rows = [list(row) for row in workbook[selected_sheet].iter_rows(values_only=True)]
                finally:
                    workbook.close()
            elif source.suffix.lower() == ".xls":
                import xlrd

                workbook = xlrd.open_workbook(str(source), on_demand=True)
                try:
                    selected_sheet = sheet_name or workbook.sheet_names()[0]
                    worksheet = workbook.sheet_by_name(selected_sheet)
                    rows = [worksheet.row_values(index) for index in range(worksheet.nrows)]
                finally:
                    release = getattr(workbook, "release_resources", None)
                    if callable(release):
                        release()
            elif source.suffix.lower() in {".csv", ".tsv", ".txt"}:
                delimiter = "\t" if source.suffix.lower() == ".tsv" else ","
                with source.open("r", encoding="utf-8-sig", newline="") as handle:
                    rows = [
                        [_coerce_delimited_cell(value) for value in row]
                        for row in csv.reader(handle, delimiter=delimiter)
                    ]
            else:
                return ResultEnvelope.fail("UNSUPPORTED_INPUT", f"Unsupported data file: {source.suffix}")
        except Exception as exc:
            return ResultEnvelope.fail("INPUT_PARSE_FAILED", str(exc))
        column_labels, data_rows = _split_header(rows, has_header)
        target_name = worksheet_name or source.stem

        rectangular, width = _rectangular_values(data_rows)

        def import_rows() -> ResultEnvelope:
            app = self._require_app()
            created_name: str | None = None
            page = None
            template_name = None
            system_template = False
            if target_mode == "new_workbook":
                template_name, system_template = _system_worksheet_template(app)
                created_name = app.CreatePage(
                    ORIGIN_PAGE_WORKSHEET,
                    target_name,
                    template_name,
                    ORIGIN_CREATE_HIDDEN,
                )
                sheet = app.FindWorksheet(created_name)
                worksheet_pages = _safe_attr(app, "WorksheetPages")
                if worksheet_pages is not None:
                    try:
                        page = worksheet_pages.Item(created_name)
                    except Exception:
                        page = None
            else:
                sheet = _resolve_worksheet(app, target_name)
            if sheet is None:
                raise LookupError(f"Origin worksheet was not created: {target_name}")

            sheet_name_actual = str(_safe_attr(sheet, "Name", "Sheet1"))
            page_name = str(created_name or target_name)
            worksheet_ref = f"[{page_name}]{sheet_name_actual}" if created_name else target_name
            template_reset = {
                "applied": target_mode == "new_workbook",
                "initial_rows": int(_safe_attr(sheet, "Rows", 0) or 0),
                "initial_columns": int(_safe_attr(sheet, "Cols", 0) or 0),
                "initial_sheet_long_name": str(_safe_attr(sheet, "LongName", "")),
                "initial_page_long_name": str(_safe_attr(page, "LongName", "")) if page else None,
                "template_name": template_name,
                "system_template": system_template,
                "data_connection_reset": (
                    "system_installation_template_bypass"
                    if system_template
                    else "fallback_template_with_strict_readback"
                ),
            }

            labels_applied = 0
            source_columns = [
                [item[index] for item in rectangular]
                for index in range(width)
            ]
            column_formats: list[int | None] = []
            self._current_stage = "write"
            try:
                set_result = sheet.SetData(rectangular, 0, 0)
            except Exception as exc:
                return ResultEnvelope.fail(
                    "IMPORT_WRITE_FAILED",
                    f"Origin worksheet block write failed: {exc}",
                    data={
                        "worksheet": target_name,
                        "worksheet_ref": worksheet_ref,
                        "stage": "write",
                        "template_reset": template_reset,
                    },
                )
            set_data_returns: list[Any] = [set_result]
            if set_result is False or (
                isinstance(set_result, int)
                and not isinstance(set_result, bool)
                and set_result == 0
            ):
                return ResultEnvelope.fail(
                    "IMPORT_WRITE_REJECTED",
                    "Origin explicitly rejected Worksheet.SetData during import",
                    data={
                        "worksheet": target_name,
                        "worksheet_ref": worksheet_ref,
                        "stage": "write",
                        "set_data_returns": set_data_returns,
                        "template_reset": template_reset,
                    },
                )

            # Apply the exact final dimensions and metadata only after the
            # successful block write. Origin can lock or invalidate a fresh
            # sheet's storage when formats/dimensions are changed first.
            if target_mode == "new_workbook":
                sheet.Cols = width
                sheet.Rows = len(rectangular)
                try:
                    sheet.LongName = ""
                except Exception:
                    pass
                if page is not None:
                    try:
                        page.LongName = target_name
                    except Exception:
                        pass

            for index, column_values in enumerate(source_columns):
                target_column = _worksheet_column(sheet, index)
                try:
                    column_formats.append(
                        _set_column_data_format(target_column, column_values, force=True)
                    )
                except Exception as exc:
                    return ResultEnvelope.fail(
                        "IMPORT_COLUMN_FORMAT_FAILED",
                        f"Origin rejected format for imported column {index}: {exc}",
                        data={
                            "worksheet": target_name,
                            "worksheet_ref": worksheet_ref,
                            "stage": "format",
                            "failed_column": index,
                            "template_reset": template_reset,
                        },
                    )
                for attribute in ("Units", "Comments"):
                    if target_mode == "new_workbook":
                        try:
                            setattr(target_column, attribute, "")
                        except Exception:
                            pass
                if column_labels and index < len(column_labels) and column_labels[index]:
                    try:
                        target_column.LongName = column_labels[index]
                        labels_applied += 1
                    except Exception:
                        pass

            self._current_stage = "validate"
            profiles: list[dict[str, Any]] = []
            mismatched_columns: list[int] = []
            empty_destination_columns: list[int] = []
            for index, source_values in enumerate(source_columns):
                destination_values = _column_variant_values(
                    _worksheet_column(sheet, index),
                    0,
                    len(rectangular) - 1,
                )
                destination_values = (
                    destination_values[: len(rectangular)]
                    + [None] * max(0, len(rectangular) - len(destination_values))
                )
                source_profile = _column_profile(source_values)
                destination_profile = _column_profile(destination_values)
                value_mismatches = [
                    row_index
                    for row_index, (expected, actual) in enumerate(
                        zip(source_values, destination_values, strict=True)
                    )
                    if not _cells_equal(expected, actual)
                ]
                if value_mismatches:
                    mismatched_columns.append(index)
                if (
                    source_profile["non_empty_count"] > 0
                    and destination_profile["non_empty_count"] == 0
                ):
                    empty_destination_columns.append(index)
                non_empty_values = [
                    _normalize_cell(value)
                    for value in destination_values
                    if not _is_missing_cell(value)
                ]
                profiles.append(
                    {
                        "index": index,
                        "name": str(_safe_attr(_worksheet_column(sheet, index), "Name", "")),
                        "label": (
                            column_labels[index]
                            if column_labels and index < len(column_labels)
                            else ""
                        ),
                        "data_format": column_formats[index],
                        "source": source_profile,
                        "destination": destination_profile,
                        "first_value": non_empty_values[0] if non_empty_values else None,
                        "last_value": non_empty_values[-1] if non_empty_values else None,
                        "mismatch_rows": value_mismatches[:10],
                    }
                )

            result_data = {
                "worksheet": target_name,
                "worksheet_ref": worksheet_ref,
                "workbook": page_name if created_name else None,
                "worksheet_name": sheet_name_actual,
                "rows": len(rectangular),
                "columns": width,
                "column_labels": column_labels,
                "column_labels_applied": labels_applied,
                "header_mode": "explicit" if has_header is not None else "auto",
                "target_mode": target_mode,
                "set_data_returns": set_data_returns,
                "column_profiles": profiles,
                "non_empty_count": sum(
                    item["destination"]["non_empty_count"] for item in profiles
                ),
                "template_reset": template_reset,
                "stage": "validate",
            }
            if empty_destination_columns:
                return ResultEnvelope.fail(
                    "IMPORT_DATA_LOSS",
                    "One or more populated source columns became empty in Origin",
                    data={
                        **result_data,
                        "failed_columns": empty_destination_columns,
                    },
                )
            if mismatched_columns:
                return ResultEnvelope.fail(
                    "IMPORT_VALIDATION_FAILED",
                    "Origin readback did not match one or more imported columns",
                    data={**result_data, "failed_columns": mismatched_columns},
                )
            return ResultEnvelope.ok(result_data)

        result = self._submit(import_rows)
        if result.success:
            warnings = []
            labels = result.data.get("column_labels") or []
            if labels and result.data.get("column_labels_applied", 0) < sum(bool(label) for label in labels):
                warnings.append("One or more imported column labels could not be applied in Origin")
            result = ResultEnvelope.ok(
                result.data,
                warnings=warnings,
                artifacts=[Artifact(str(source), "input_data")],
                duration_ms=result.duration_ms,
            )
        return result

    @_serialized_operation
    def run_xfunction(
        self,
        *,
        name: str,
        parameters: Mapping[str, Any],
        outputs: Mapping[str, Any] | None = None,
        create_operation: bool = False,
        recalculate_mode: str = "none",
        allow_unverified: bool = False,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        try:
            plan = build_xfunction_plan(
                name,
                parameters,
                outputs=outputs,
                create_operation=create_operation,
                recalculate_mode=recalculate_mode,
                allow_unverified=allow_unverified,
            )
        except NativeValidationError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))

        def execute() -> ResultEnvelope:
            try:
                return ResultEnvelope.ok(
                    execute_xfunction_plan(
                        self._require_app(), plan, self._analysis_operations
                    )
                )
            except NativeValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))

        return self._submit(execute, stage="run_xfunction")

    @_serialized_operation
    def list_analysis_operations(self, *, scope_ref: str | None = None) -> ResultEnvelope:
        def collect() -> ResultEnvelope:
            self._require_app()
            try:
                operations = self._analysis_operations.list(scope_ref)
            except NativeValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))
            return ResultEnvelope.ok(
                {
                    "operations": operations,
                    "count": len(operations),
                    "scope_ref": scope_ref,
                    "managed_only": True,
                }
            )

        return self._submit(
            collect, retryable=True, stage="list_analysis_operations"
        )

    @_serialized_operation
    def get_analysis_operation(self, *, operation_ref: str) -> ResultEnvelope:
        def read() -> ResultEnvelope:
            try:
                return ResultEnvelope.ok(
                    read_analysis_operation(
                        self._require_app(), operation_ref, self._analysis_operations
                    )
                )
            except NativeValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))

        return self._submit(read, retryable=True, stage="get_analysis_operation")

    @_serialized_operation
    def recalculate_analysis(
        self, *, operation_ref: str, wait: bool = True
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard

        def recalculate() -> ResultEnvelope:
            try:
                return ResultEnvelope.ok(
                    recalculate_operation(
                        self._require_app(),
                        operation_ref,
                        self._analysis_operations,
                        wait=wait,
                    )
                )
            except NativeValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))

        return self._submit(recalculate, stage="recalculate_analysis")

    @_serialized_operation
    def manage_analysis_template(
        self,
        *,
        action: str,
        path: str,
        workbook_ref: str | None = None,
        overwrite: bool = False,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        try:
            plan = build_analysis_template_plan(
                action=action,
                path=path,
                workbook_ref=workbook_ref,
                overwrite=overwrite,
            )
        except NativeValidationError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))

        def execute() -> ResultEnvelope:
            try:
                data = execute_analysis_template_plan(self._require_app(), plan)
            except NativeValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))
            artifacts = (
                [Artifact(str(plan.path), "analysis_template")]
                if plan.action == "save"
                else []
            )
            return ResultEnvelope.ok(data, artifacts=artifacts)

        return self._submit(execute, stage=f"analysis_template_{plan.action}")

    @_serialized_operation
    def execute_labtalk(
        self,
        *,
        script: str,
        segments: list[str] | None = None,
        result_variable: str | None = None,
        result_variables: list[str] | None = None,
        result_numeric_variables: list[str] | None = None,
        result_string_variables: list[str] | None = None,
        warning_variable: str | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        if segments is not None and script.strip():
            return ResultEnvelope.fail(
                "AMBIGUOUS_LABTALK_INPUT",
                "Provide either script or explicit segments, not both",
            )
        commands = list(segments) if segments is not None else [script]
        if not commands or any(not command.strip() for command in commands):
            return ResultEnvelope.fail(
                "EMPTY_LABTALK",
                "LabTalk script or every explicit segment must be non-empty",
            )
        numeric_names: list[str] = []
        for name in [result_variable, *(result_variables or []), *(result_numeric_variables or [])]:
            if name and name not in numeric_names:
                numeric_names.append(name)
        string_names: list[str] = []
        for name in result_string_variables or []:
            if name and name not in string_names:
                string_names.append(name)
        overlap = sorted(set(numeric_names) & set(string_names))
        if overlap:
            return ResultEnvelope.fail(
                "AMBIGUOUS_LABTALK_RESULT",
                f"Variables cannot be declared as both numeric and string: {', '.join(overlap)}",
            )

        def execute() -> ResultEnvelope:
            app = self._require_app()
            for segment_index, command in enumerate(commands, start=1):
                self._current_stage = "execute_segment"
                executed = app.Execute(command)
                if executed is False:
                    return ResultEnvelope.fail(
                        "LABTALK_EXECUTION_FAILED",
                        f"Origin rejected LabTalk segment {segment_index}",
                        data={
                            "stage": "execute_segment",
                            "failed_segment": segment_index,
                            "failed_statement": command,
                            "completed_segments": segment_index - 1,
                            "total_segments": len(commands),
                            "origin_error_numeric": _safe_call(
                                app, "LTVar", "@LSE", default=None
                            ),
                            "origin_error_string": _safe_call(
                                app, "LTStr", "%Z", default=None
                            ),
                        },
                    )

            self._current_stage = "collect_results"
            numeric_output = {name: app.LTVar(name) for name in numeric_names}
            string_output = {name: app.LTStr(name) for name in string_names}
            output = {**numeric_output, **string_output}
            warnings: list[str] = []
            if result_variable or result_variables:
                warnings.append(
                    "result_variable/result_variables are deprecated numeric-only aliases; "
                    "use result_numeric_variables or result_string_variables"
                )
            if warning_variable:
                warning_value = app.LTStr(warning_variable)
                warnings.extend(
                    line.strip()
                    for line in str(warning_value).splitlines()
                    if line.strip()
                )
            return ResultEnvelope.ok(
                {
                    "executed": True,
                    "segmented": segments is not None,
                    "completed_segments": len(commands),
                    "total_segments": len(commands),
                    "result_variable": result_variable,
                    "result": output.get(result_variable) if result_variable else None,
                    "output": output,
                    "numeric_output": numeric_output,
                    "string_output": string_output,
                    "warning_variable": warning_variable,
                },
                warnings=warnings,
            )

        return self._submit(execute, stage="execute_segment")

    @_serialized_operation
    def run_analysis(
        self,
        *,
        worksheet_name: str,
        method: str,
        x_column: str | int,
        y_column: str | int,
        options: dict[str, Any] | None = None,
        row_start: int = 0,
        row_end: int = -1,
        filters: list[Mapping[str, Any]] | None = None,
        row_order: str = "as_is",
    ) -> ResultEnvelope:
        try:
            request = build_analysis_request(
                method=method,
                x_column=str(x_column),
                y_column=str(y_column),
                options=options,
            )
        except ValueError as exc:
            return ResultEnvelope.fail("INVALID_ANALYSIS_REQUEST", str(exc))
        if row_start < 0 or (row_end != -1 and row_end < row_start):
            return ResultEnvelope.fail(
                "INVALID_ANALYSIS_SELECTION",
                "row_start must be non-negative and row_end must be -1 or at least row_start",
            )
        if row_order not in {"as_is", "reverse"}:
            return ResultEnvelope.fail(
                "INVALID_ANALYSIS_SELECTION",
                "row_order must be as_is or reverse",
            )
        try:
            normalized_filters = _normalize_analysis_filters(filters)
        except (KeyError, TypeError, ValueError) as exc:
            return ResultEnvelope.fail("INVALID_ANALYSIS_SELECTION", str(exc))

        backend = str(request.options.get("backend", "python"))
        if backend == "origin_native":
            guard = self._guard_mutation()
            if guard:
                return guard
            if row_start != 0 or row_end != -1 or normalized_filters or row_order != "as_is":
                return ResultEnvelope.fail(
                    "NATIVE_ANALYSIS_SELECTION_UNSUPPORTED",
                    "Origin-native analysis currently requires the full explicit columns with no filters or row reordering",
                )
            native_options = {
                key: value
                for key, value in request.options.items()
                if key not in {"backend", "create_operation", "recalculate_mode"}
            }
            create_operation = bool(request.options.get("create_operation", False))
            recalculate_mode = str(request.options.get("recalculate_mode", "none"))
            worksheet_ref = RangeRef(worksheet_name).value
            try:
                if method == "linear_fit" and not native_options:
                    plan = build_xfunction_plan(
                        "fitlr",
                        {
                            "iy": RangeRef(
                                f"{worksheet_ref}!({x_column},{y_column})"
                            )
                        },
                        create_operation=create_operation,
                        recalculate_mode=recalculate_mode,
                    )
                elif method == "fft" and set(native_options) <= {"sample_spacing"}:
                    if native_options.get("sample_spacing", 1.0) != 1.0:
                        return ResultEnvelope.fail(
                            "NATIVE_ANALYSIS_OPTION_UNSUPPORTED",
                            "Origin-native FFT sample_spacing mapping is not yet verified",
                        )
                    plan = build_xfunction_plan(
                        "fft1",
                        {
                            "ix": RangeRef(f"{worksheet_ref}!{x_column}"),
                            "iy": RangeRef(f"{worksheet_ref}!{y_column}"),
                        },
                        create_operation=create_operation,
                        recalculate_mode=recalculate_mode,
                    )
                else:
                    return ResultEnvelope.fail(
                        "NATIVE_ANALYSIS_UNSUPPORTED",
                        f"No verified Origin-native mapping is available for {method} with these options",
                    )
            except NativeValidationError as exc:
                return ResultEnvelope.fail(exc.code, str(exc))

            def execute_native() -> ResultEnvelope:
                try:
                    data = execute_xfunction_plan(
                        self._require_app(), plan, self._analysis_operations
                    )
                except NativeValidationError as exc:
                    return ResultEnvelope.fail(exc.code, str(exc))
                return ResultEnvelope.ok({**data, "backend": "origin_native"})

            return self._submit(execute_native, stage=f"native_analysis_{method}")

        def analyze() -> dict[str, Any]:
            app = self._require_app()
            sheet = _resolve_worksheet(app, worksheet_name)
            x_index = _column_index(sheet, x_column)
            y_index = _column_index(sheet, y_column)
            raw = table_from_com_value(
                sheet.GetData(row_start, 0, row_end, -1, ORIGIN_ARRAY2D_NUMERIC)
            )
            x_values: list[Any] = []
            y_values: list[Any] = []
            for row in raw:
                if max(x_index, y_index) >= len(row):
                    continue
                if row[x_index] is None or row[y_index] is None:
                    continue
                x_value = float(row[x_index])
                y_value = float(row[y_index])
                values = {"x": x_value, "y": y_value}
                if not all(
                    FILTER_OPERATORS[condition["operator"]](
                        values[condition["column"]], condition["value"]
                    )
                    for condition in normalized_filters
                ):
                    continue
                x_values.append(x_value)
                y_values.append(y_value)
            if row_order == "reverse":
                x_values.reverse()
                y_values.reverse()
            try:
                analysis_result = run_analysis_data(method, x_values, y_values, options)
            except (TypeError, ValueError) as exc:
                raise AnalysisExecutionError(str(exc)) from exc
            return {
                "worksheet": worksheet_name,
                "method": method,
                "x_column": x_column,
                "y_column": y_column,
                "selection": {
                    "row_start": row_start,
                    "row_end": row_end,
                    "row_order": row_order,
                    "filters": normalized_filters,
                    "rows_read": len(raw),
                    "rows_after_filter": len(x_values),
                },
                "result": analysis_result,
            }

        return self._submit(analyze, retryable=True)

    @_serialized_operation
    def create_graph(
        self,
        *,
        graph_type: str,
        roles: Mapping[str, Any],
        graph_name: str | None = None,
        allow_unverified: bool = False,
    ) -> ResultEnvelope:
        try:
            catalog_entry = validate_graph_request(graph_type, roles)
        except GraphCatalogError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))
        if catalog_entry["status"] == "verified":
            if not roles.get("worksheet"):
                return ResultEnvelope.fail(
                    "GRAPH_WORKSHEET_REQUIRED",
                    "Verified worksheet graphs require the worksheet role",
                )
            y_value = roles["y"]
            y_columns = list(y_value) if isinstance(y_value, (list, tuple)) else [y_value]
            return self.create_plot(
                worksheet_name=str(roles.get("worksheet", "")),
                graph_type=graph_type,
                x_column=roles["x"],
                y_columns=y_columns,
                label_column=roles.get("label"),
                graph_name=graph_name,
            )
        if not allow_unverified:
            return ResultEnvelope.fail(
                "GRAPH_TYPE_UNVERIFIED",
                f"{graph_type} is supported-unverified; set allow_unverified explicitly",
                data={"catalog": catalog_entry},
            )
        guard = self._guard_mutation()
        if guard:
            return guard
        try:
            serialized_roles = {
                key: RangeRef(str(value)).value for key, value in roles.items()
            }
        except NativeValidationError as exc:
            return ResultEnvelope.fail(exc.code, str(exc))
        plot_id = catalog_entry["plot_id"]
        if plot_id is None:
            return ResultEnvelope.fail(
                "GRAPH_EXECUTION_MAPPING_UNVERIFIED",
                f"No numeric Origin plot mapping is enabled for {graph_type}",
                data={"catalog": catalog_entry},
            )
        if catalog_entry["input_kind"] == "matrix":
            command = f"plotm im:={serialized_roles['z']} plot:={plot_id};"
        else:
            ordered = [serialized_roles[key] for key in catalog_entry["required_roles"]]
            command = f"plotxy iy:=({','.join(ordered)}) plot:={plot_id};"

        def execute() -> ResultEnvelope:
            result = self._require_app().Execute(command)
            if result is False or result == 0:
                return ResultEnvelope.fail(
                    "GRAPH_CREATE_REJECTED", f"Origin rejected {graph_type} creation"
                )
            return ResultEnvelope.ok(
                {
                    "graph_type": graph_type,
                    "graph_name": graph_name,
                    "roles": serialized_roles,
                    "status": "supported_unverified",
                    "origin_command_accepted": True,
                },
                warnings=["Graph creation is not yet covered by a live version-specific smoke test"],
            )

        return self._submit(execute, stage=f"create_graph_{graph_type}")

    @_serialized_operation
    def manage_graph_layout(
        self,
        *,
        action: str,
        graph_ref: str,
        source_graph_refs: list[str] | None = None,
        layer_refs: list[str] | None = None,
        position: list[float] | None = None,
        rows: int | None = None,
        columns: int | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        try:
            plan = build_layout_plan(
                action=action,
                graph_ref=graph_ref,
                source_graph_refs=source_graph_refs,
                layer_refs=layer_refs,
                position=position,
                rows=rows,
                columns=columns,
            )
        except (GraphLayoutError, ObjectPlanError) as exc:
            return ResultEnvelope.fail(getattr(exc, "code", "GRAPH_LAYOUT_INVALID"), str(exc))

        def execute() -> ResultEnvelope:
            app = self._require_app()
            layer = _safe_call(app, "FindGraphLayer", plan.graph_ref, default=None)
            if layer is None:
                layer = _safe_call(app, "FindGraphLayer", f"[{plan.graph_ref}]1", default=None)
            if layer is None:
                return ResultEnvelope.fail("GRAPH_NOT_FOUND", f"Graph not found: {plan.graph_ref}")
            result = layer.Execute(plan.command)
            if result is False or result == 0:
                return ResultEnvelope.fail("GRAPH_LAYOUT_REJECTED", f"Origin rejected layout {plan.action}")
            return ResultEnvelope.ok(
                {
                    "action": plan.action,
                    "graph_ref": plan.graph_ref,
                    "source_graph_refs": list(plan.source_graph_refs),
                    "position": list(plan.position) if plan.position else None,
                    "expected_layer_delta": plan.expected_layer_delta,
                    "status": "supported_unverified",
                }
            )

        return self._submit(execute, stage=f"graph_layout_{plan.action}")

    @_serialized_operation
    def apply_graph_template(
        self,
        *,
        graph_ref: str,
        template_path: str,
        expected_sha256: str,
        required_layers: int | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard

        def execute() -> ResultEnvelope:
            app = self._require_app()
            page = _safe_call(_safe_attr(app, "GraphPages"), "Item", graph_ref, default=None)
            if page is None:
                return ResultEnvelope.fail("GRAPH_NOT_FOUND", f"Graph not found: {graph_ref}")
            layers = _safe_attr(page, "Layers")
            actual_layers = int(_safe_attr(layers, "Count", 0) or 0)
            try:
                plan = apply_template_plan(
                    graph_ref=graph_ref,
                    template_path=template_path,
                    expected_sha256=expected_sha256,
                    required_layers=required_layers,
                    actual_layers=actual_layers,
                )
            except (GraphTemplateError, ObjectPlanError) as exc:
                return ResultEnvelope.fail(getattr(exc, "code", "GRAPH_TEMPLATE_INVALID"), str(exc))
            executor = getattr(page, "Execute", None)
            if not callable(executor) and actual_layers:
                executor = getattr(layers.Item(0), "Execute", None)
            if not callable(executor):
                return ResultEnvelope.fail("GRAPH_TEMPLATE_INTERFACE_UNAVAILABLE", "Graph does not expose Execute")
            result = executor(plan.command)
            if result is False or result == 0:
                return ResultEnvelope.fail("GRAPH_TEMPLATE_REJECTED", "Origin rejected template application")
            return ResultEnvelope.ok(
                {
                    "graph_ref": graph_ref,
                    "template_path": str(plan.template_path),
                    "sha256": plan.sha256,
                    "layers": actual_layers,
                    "status": "supported_unverified",
                }
            )

        return self._submit(execute, stage="apply_graph_template")

    def view_graph(
        self,
        *,
        graph_name: str,
        output_path: str | None = None,
        expected_colors: list[str] | None = None,
        tolerance: int = 12,
    ) -> ResultEnvelope:
        target = Path(output_path).expanduser().resolve() if output_path else (
            Path(tempfile.gettempdir()) / f"origin-preview-{uuid4().hex}.png"
        )
        exported = self.export_graph(
            graph_name=graph_name,
            output_path=str(target),
            export_format="png",
            overwrite="replace",
        )
        if not exported.success:
            return exported
        try:
            metrics = inspect_png(target, expected_colors=expected_colors, tolerance=tolerance)
        except (OSError, ValueError) as exc:
            return ResultEnvelope.fail(
                "GRAPH_PREVIEW_INVALID",
                str(exc),
                artifacts=exported.artifacts,
            )
        return ResultEnvelope.ok(
            {
                **(exported.data or {}),
                "preview_path": str(target),
                "pixel_metrics": metrics,
            },
            artifacts=exported.artifacts,
            warnings=exported.warnings,
            duration_ms=exported.duration_ms,
            origin_version=exported.origin_version,
        )

    @_serialized_operation
    def create_plot(
        self,
        *,
        worksheet_name: str,
        graph_type: str,
        x_column: str | int,
        y_columns: list[str | int],
        label_column: str | int | None = None,
        graph_name: str | None = None,
        options: dict[str, Any] | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        try:
            spec = build_graph_spec(
                graph_type=graph_type,
                x_column=str(x_column),
                y_columns=[str(value) for value in y_columns],
                options=options,
            )
        except ValueError as exc:
            return ResultEnvelope.fail("INVALID_GRAPH_REQUEST", str(exc))
        if label_column is not None and spec.graph_type != "multi_layer" and len(y_columns) != 1:
            return ResultEnvelope.fail(
                "INVALID_GRAPH_REQUEST",
                "label_column requires one Y column unless graph_type is multi_layer",
            )

        def create() -> dict[str, Any]:
            app = self._require_app()
            sheet = _resolve_worksheet(app, worksheet_name)
            x_index = _column_index(sheet, x_column)
            y_indices = [_column_index(sheet, value) for value in y_columns]
            label_column_obj = (
                _worksheet_column(sheet, label_column) if label_column is not None else None
            )
            template = str(spec.options.get("template", "Origin"))
            created_name = app.CreatePage(
                ORIGIN_PAGE_GRAPH,
                graph_name or "",
                template,
                ORIGIN_CREATE_HIDDEN,
            )
            graph_layer = app.FindGraphLayer(created_name)
            if graph_layer is None:
                raise LookupError(f"Origin graph layer was not created: {created_name}")
            plot_type = {
                "line": ORIGIN_PLOT_LINE,
                "scatter": ORIGIN_PLOT_SCATTER,
                "line_symbol": ORIGIN_PLOT_LINESYMB,
                "semilog": ORIGIN_PLOT_LINE,
                "loglog": ORIGIN_PLOT_LINE,
                "bar": ORIGIN_PLOT_COLUMN,
                "multi_layer": ORIGIN_PLOT_LINESYMB,
            }[spec.graph_type]
            layers_created = 1
            if spec.graph_type == "multi_layer":
                page = app.GraphPages.Item(created_name)
                layers = page.Layers
                layers_created = 0
                for position, y_index in enumerate(y_indices):
                    if position == 0:
                        layer = graph_layer
                    else:
                        layers.Add()
                        layer = app.FindGraphLayer(f"[{created_name}]{position + 1}")
                        if layer is None:
                            raise LookupError(
                                f"Origin graph layer was not created: [{created_name}]{position + 1}"
                            )
                    data_range = app.NewDataRange()
                    data_range.Add("X", sheet, 0, x_index, -1, x_index)
                    data_range.Add("Y", sheet, 0, y_index, -1, y_index)
                    plot = layer.DataPlots.Add(data_range, plot_type)
                    if plot is None:
                        raise OSError("Origin did not create the data plot")
                    if label_column_obj is not None:
                        _execute_graph_commands(
                            layer,
                            _data_plot_label_commands(plot, label_column_obj),
                            "Origin rejected the plot label binding",
                        )
                    layer.Execute("rescale;")
                    layers_created += 1
            else:
                data_range = app.NewDataRange()
                data_range.Add("X", sheet, 0, x_index, -1, x_index)
                for index in y_indices:
                    data_range.Add("Y", sheet, 0, index, -1, index)
                plot = graph_layer.DataPlots.Add(data_range, plot_type)
                if plot is None:
                    raise OSError("Origin did not create the data plot")
                if label_column_obj is not None:
                    _execute_graph_commands(
                        graph_layer,
                        _data_plot_label_commands(plot, label_column_obj),
                        "Origin rejected the plot label binding",
                    )
                commands = ["rescale;"]
                if spec.graph_type in {"semilog", "loglog"}:
                    commands.append(f"layer.y.type={SCALE_TYPES['log10']};")
                if spec.graph_type == "loglog":
                    commands.append(f"layer.x.type={SCALE_TYPES['log10']};")
                graph_layer.Execute("".join(commands))
            return {
                "graph_name": str(created_name),
                "graph_id": f"[{created_name}]1",
                "worksheet": worksheet_name,
                "graph_type": spec.graph_type,
                "x_column": x_column,
                "y_columns": y_columns,
                "label_column": label_column,
                "layers_created": layers_created,
            }

        return self._submit(create)

    @_serialized_operation
    def configure_graph(
        self,
        *,
        graph_name: str,
        options: Mapping[str, Any],
        labtalk: str | None = None,
    ) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        try:
            commands = _graph_configuration_commands(options, labtalk)
            data_binding = _normalize_graph_data_binding(options.get("data_binding"))
            categorical_style = _normalize_categorical_style(options.get("categorical_style"))
            categorical_legend = _normalize_categorical_legend(options.get("legend"))
            if categorical_legend and not categorical_style:
                raise ValueError(
                    "A categorical legend sourced from plot_style_mapping requires categorical_style"
                )
            if (
                categorical_legend
                and categorical_style
                and categorical_legend["plot_index"] != categorical_style["plot_index"]
            ):
                raise ValueError("legend and categorical_style plot_index must match")
        except (TypeError, ValueError) as exc:
            return ResultEnvelope.fail("INVALID_GRAPH_OPTIONS", str(exc))

        def configure_core() -> dict[str, Any]:
            app = self._require_app()
            layer = _resolve_graph_layer(app, graph_name)
            binding_commands: list[str] = []
            if data_binding:
                sheet = _resolve_worksheet(app, data_binding["worksheet_name"])
                x_index = _column_index(sheet, data_binding["x_column"])
                y_indices = [
                    _column_index(sheet, column) for column in data_binding["y_columns"]
                ]
                label_column_obj = (
                    _worksheet_column(sheet, data_binding["label_column"])
                    if data_binding.get("label_column") is not None
                    else None
                )
                old_plots = _collection_items(layer.DataPlots)
                new_plots: list[Any] = []
                try:
                    for y_index in y_indices:
                        data_range = app.NewDataRange()
                        data_range.Add("X", sheet, 0, x_index, -1, x_index)
                        data_range.Add("Y", sheet, 0, y_index, -1, y_index)
                        plot = layer.DataPlots.Add(
                            data_range,
                            DATA_BINDING_PLOT_TYPES[data_binding["plot_type"]],
                        )
                        if plot is None:
                            raise OSError("Origin did not create the replacement data plot")
                        new_plots.append(plot)
                        if label_column_obj is not None:
                            binding_commands.extend(
                                _data_plot_label_commands(plot, label_column_obj)
                            )
                    _execute_graph_commands(
                        layer,
                        binding_commands,
                        "Origin rejected the replacement plot label binding",
                    )
                except Exception:
                    for plot in new_plots:
                        if any(_same_origin_object(plot, old) for old in old_plots):
                            continue
                        try:
                            plot.Destroy()
                        except Exception:
                            pass
                    raise
                for plot in old_plots:
                    if any(_same_origin_object(plot, new) for new in new_plots):
                        continue
                    if plot.Destroy() is False:
                        raise OSError("Origin could not remove an old data plot after rebinding")
            applied_commands = list(commands)
            categorical_result = None
            if categorical_style:
                category_commands, categorical_result = _categorical_style_commands(
                    app,
                    layer,
                    categorical_style,
                    data_binding,
                )
                applied_commands.extend(category_commands)
            _execute_graph_commands(
                layer,
                applied_commands,
                "Origin rejected the graph configuration commands",
            )
            return {
                "graph_name": graph_name,
                "options_applied": sorted(options),
                "commands_applied": len(binding_commands) + len(applied_commands),
                "data_binding": data_binding,
                "categorical_style": categorical_result,
                "legend": categorical_legend,
            }

        configured = self._submit(configure_core)
        if not configured.success or not categorical_legend:
            return configured

        def configure_legend() -> dict[str, Any]:
            app = self._require_app()
            layer = _resolve_graph_layer(app, graph_name)
            if categorical_legend["replace_existing"]:
                for object_name in ("CategoryKey", "Legend"):
                    try:
                        # Origin can report failure after actually removing a label.
                        # Verify object state below instead of trusting this return value.
                        layer.Execute(f"label -r {object_name};")
                    except Exception:
                        pass
            categorical_result = configured.data["categorical_style"]
            legend_commands = _categorical_legend_commands(
                categorical_legend,
                categorical_result,
            )
            _execute_graph_commands(
                layer,
                legend_commands[:1],
                "Origin rejected categorical Legend creation",
            )
            _execute_graph_commands(
                layer,
                legend_commands[1:],
                "Origin rejected categorical Legend formatting",
            )
            expected_entries = len(categorical_result["category_counts"])
            if categorical_legend["show_all_categories"]:
                expected_entries += len(categorical_result["unused_categories"])
            _verify_categorical_legend(
                app,
                layer,
                categorical_legend["replace_existing"],
                expected_entries=expected_entries,
            )
            return {"commands_applied": len(legend_commands)}

        legend_result = self._submit(configure_legend)
        if not legend_result.success:
            return replace(
                legend_result,
                data={
                    "graph_name": graph_name,
                    "configuration_applied": True,
                    "legend_verified": False,
                },
                warnings=[
                    *legend_result.warnings,
                    "Graph configuration was applied before categorical Legend verification failed",
                ],
                duration_ms=(configured.duration_ms or 0) + (legend_result.duration_ms or 0),
            )
        return ResultEnvelope.ok(
            {
                **configured.data,
                "commands_applied": configured.data["commands_applied"]
                + legend_result.data["commands_applied"],
            },
            warnings=[*configured.warnings, *legend_result.warnings],
            duration_ms=(configured.duration_ms or 0) + (legend_result.duration_ms or 0),
        )

    @_serialized_operation
    def export_graph(
        self,
        *,
        graph_name: str,
        output_path: str,
        export_format: str,
        overwrite: str = "skip",
    ) -> ResultEnvelope:
        try:
            fmt = normalize_export_format(export_format)
        except ValueError as exc:
            return ResultEnvelope.fail("UNSUPPORTED_EXPORT", str(exc))
        target = Path(output_path).expanduser().resolve()
        if target.exists() and overwrite == "skip":
            return ResultEnvelope.fail("OUTPUT_EXISTS", f"Refusing to overwrite existing graph export: {target}")
        if overwrite not in {"skip", "rename", "replace"}:
            return ResultEnvelope.fail("INVALID_OVERWRITE", "overwrite must be skip, rename, or replace")

        def export() -> dict[str, Any]:
            app = self._require_app()
            _resolve_graph_layer(app, graph_name)
            target.parent.mkdir(parents=True, exist_ok=True)
            before = snapshot_artifacts(target.parent, target.stem)
            script = (
                f"expGraph export:=specified pages:={_labtalk_quote(_graph_page_name(graph_name))} "
                f"type:={fmt} path:={_labtalk_quote(target.parent)} "
                f"filename:={_labtalk_quote(target.stem)} overwrite:={overwrite};"
            )
            if app.Execute(script) is False:
                raise LabTalkExecutionError("Origin graph export command failed")
            actual = select_changed_artifact(
                target.parent,
                target.stem,
                before,
                suffixes=export_suffixes(fmt),
            )
            if actual is None:
                raise FileNotFoundError(f"Origin did not create or change an export artifact: {target}")
            return {"graph_name": graph_name, "path": str(actual), "format": fmt}

        result = self._submit(export)
        if result.success:
            path = result.data["path"]
            try:
                actual = validate_export_path(path, kind=fmt)
                result = ResultEnvelope.ok(
                    result.data,
                    artifacts=[Artifact(str(actual), "graph_export")],
                    duration_ms=result.duration_ms,
                )
            except (OSError, ValueError) as exc:
                return ResultEnvelope.fail("EXPORT_VALIDATION_FAILED", str(exc))
        return result

    @_serialized_operation
    def close_project(self, *, discard_changes: bool = False) -> ResultEnvelope:
        guard = self._guard_mutation()
        if guard:
            return guard
        def close() -> dict[str, Any]:
            app = self._require_app()
            if bool(_safe_attr(app, "IsModified", False)) and not discard_changes:
                raise OriginAutomationError("Project has unsaved changes; pass discard_changes=true to close it")
            created = app.NewProject()
            if created not in (True, 1):
                raise ProjectCloseUnconfirmedError(
                    "Origin did not confirm that a new project replaced the active project"
                )
            self._project_path = None
            self._source_path = None
            return {"closed": True, "discarded_changes": discard_changes}

        return self._submit(close)

    @_serialized_operation
    def shutdown(self) -> ResultEnvelope:
        session_id = self._session_id
        if not session_id:
            return ResultEnvelope.ok({"shutdown": False, "reason": "no_active_session"})
        record = self.sessions.get(session_id)

        if self._poisoned:
            try:
                current_pids = set(self.process_snapshot())
            except Exception:
                current_pids = None
            process_still_running = (
                record.pid in current_pids
                if record.pid is not None and current_pids is not None
                else None
            )
            return ResultEnvelope.fail(
                "SESSION_POISONED",
                "Normal shutdown was skipped because the COM worker is blocked; use origin_recover_session",
                data={
                    "shutdown": False,
                    "session_id": session_id,
                    "owned": record.owned,
                    "owned_pid": record.pid,
                    "process_still_running": process_still_running,
                    "recovery_tool": "origin_recover_session",
                    "stage": "poisoned_fast_fail",
                },
                duration_ms=0,
            )

        try:
            pids_before = set(self.process_snapshot())
        except Exception:
            pids_before = None

        def close() -> dict[str, Any]:
            app = self._require_app()
            try:
                if record.owned:
                    try:
                        if self._exclusive:
                            end_session = getattr(app, "EndSession", None)
                            if callable(end_session):
                                end_session()
                    finally:
                        prompt_result = app.Execute(
                            ORIGIN_DISABLE_SAVE_PROMPT_SCRIPT
                        )
                        if prompt_result is False:
                            raise OriginAutomationError(
                                "Origin explicitly rejected save-prompt suppression"
                            )
                        new_project = app.NewProject()
                        if new_project not in (True, 1):
                            raise OriginAutomationError(
                                "Origin could not release the owned temporary project"
                            )
                        exit_result = app.Execute(ORIGIN_DELAYED_EXIT_SCRIPT)
                        if exit_result is False:
                            raise OriginAutomationError(
                                "Origin explicitly rejected delayed exit scheduling"
                            )
            finally:
                self._clear_session_state()
                app = None
            return {
                "shutdown": True,
                "owned": record.owned,
                "pid": record.pid,
                "shutdown_method": "labtalk_delayed_exit",
            }

        if record.owned:
            result = self._submit(close, allow_poisoned=True)
            record.active = False
            if result.success and record.pid is not None:
                exited, pids_after = self._wait_for_pid_exit(record.pid)
                shutdown_data = {
                    **(result.data or {}),
                    "pids_before": sorted(pids_before) if pids_before is not None else None,
                    "pids_after": sorted(pids_after) if pids_after is not None else None,
                    "exit_confirmed": exited,
                }
                if not exited:
                    result = ResultEnvelope.fail(
                        "SHUTDOWN_UNCONFIRMED",
                        f"Origin Exit returned but owned PID {record.pid} is still present or could not be checked",
                        data=shutdown_data,
                        duration_ms=result.duration_ms,
                    )
                else:
                    result = ResultEnvelope.ok(
                        shutdown_data,
                        warnings=result.warnings,
                        duration_ms=result.duration_ms,
                    )
        else:
            # Detach without invoking Exit on a user-owned Origin instance.
            def detach() -> dict[str, Any]:
                app = self._require_app()
                try:
                    if self._exclusive:
                        end_session = getattr(app, "EndSession", None)
                        if callable(end_session):
                            end_session()
                finally:
                    self._clear_session_state()
                return {"shutdown": False, "detached": True, "owned": False}

            result = self._submit(detach, allow_poisoned=True)
            self.sessions.detach(session_id)
        stop = getattr(self.worker, "stop", None)
        if callable(stop):
            try:
                stop()
            except Exception as exc:
                if result.success:
                    return ResultEnvelope.fail("WORKER_STOP_FAILED", str(exc))
        return result
